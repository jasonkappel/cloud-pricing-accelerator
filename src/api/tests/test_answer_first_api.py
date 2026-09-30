"""Tests for the additive API fields that support the answer-first web UX."""

import io
import json
from datetime import datetime, timedelta
from decimal import Decimal
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import create_app
from app import pricing
from app.data import (
    ApplicationConfigurationError,
    InMemoryApplicationRepository,
    application_max_age,
    application_max_age_hours,
)
from app.models import GapResolutionRequest
from app.pricing import PricingConfigurationError, pricing_engine

UBUNTU_EDITS = [
    ("2 Servers", "E3", "Ubuntu 22.04"),
    ("2 Servers", "F3", "Open source"),
    ("2 Servers", "E4", "Ubuntu 22.04"),
    ("2 Servers", "F4", "Open source"),
]


def _sample(repository_root: Path) -> Path:
    return repository_root / "samples" / "synthetic_intake_completed.xlsx"


def _variant(repository_root: Path, tmp_path: Path, edits: list, name: str) -> Path:
    workbook = load_workbook(_sample(repository_root))
    for sheet, cell, value in edits:
        workbook[sheet][cell] = value
    path = tmp_path / name
    workbook.save(path)
    return path


def _upload(client: TestClient, path: Path) -> dict:
    with path.open("rb") as workbook:
        response = client.post(
            "/api/intakes",
            files={
                "file": (
                    path.name,
                    workbook,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
    assert response.status_code == 201, response.text
    return response.json()


def _resolve_all(client: TestClient, detail: dict, runtime: str = "220") -> dict:
    application_id = detail["application"]["id"]
    current = detail
    for gap in detail["gaps"]:
        if gap["kind"] == "ApprovedRegions":
            body = {"azure_region": "eastus2", "aws_region": "us-east-1"}
        elif gap["kind"] == "RuntimeHours":
            body = {"runtime_hours_month": runtime}
        else:
            raise AssertionError(f"Unexpected Gap kind {gap['kind']}")
        response = client.post(
            f"/api/applications/{application_id}/gaps/{gap['id']}/resolve",
            json={"resolved_by": "Test Reviewer", **body},
        )
        assert response.status_code == 200, response.text
        current = response.json()
    return current


# ---------------------------------------------------------------- max age


@pytest.mark.parametrize(
    "value", ["0", "-1", "abc", "nan", "inf", "8761", "0.0001", "0.01", "   "]
)
def test_invalid_application_max_age_fails_startup(
    monkeypatch: pytest.MonkeyPatch,
    value: str,
) -> None:
    monkeypatch.setenv("APPLICATION_MAX_AGE_HOURS", value)
    with pytest.raises(ApplicationConfigurationError, match="APPLICATION_MAX_AGE_HOURS"):
        create_app()


def test_application_max_age_keeps_fractional_hours(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("APPLICATION_MAX_AGE_HOURS", "0.0175")
    assert application_max_age() == timedelta(seconds=63)


def test_application_max_age_defaults_to_one_hour(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.delenv("APPLICATION_MAX_AGE_HOURS", raising=False)
    assert application_max_age_hours() == Decimal("1")


def test_application_max_age_setting_controls_expiry(
    client: TestClient,
    repository_root: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setenv("APPLICATION_MAX_AGE_HOURS", "2.5")
    create_app()
    detail = _upload(client, _sample(repository_root))
    created = datetime.fromisoformat(detail["application"]["created_at"])
    assert datetime.fromisoformat(detail["expires_at"]) == created + timedelta(hours=2.5)

    listed = client.get("/api/applications").json()
    assert datetime.fromisoformat(listed[0]["expires_at"]) == created + timedelta(hours=2.5)
    assert client.get("/api/capabilities").json()["applicationMaxAgeHours"] == "2.5"


# ---------------------------------------------------------------- capabilities


@pytest.mark.parametrize(("flag", "mode"), [("true", "both"), ("false", "azure")])
def test_capabilities_report_default_mode_and_price_book(
    client: TestClient,
    monkeypatch: pytest.MonkeyPatch,
    flag: str,
    mode: str,
) -> None:
    monkeypatch.setenv("PRESENT_AWS_PRICING", flag)
    monkeypatch.delenv("APPLICATION_MAX_AGE_HOURS", raising=False)
    body = client.get("/api/capabilities").json()
    assert body["defaultMode"] == mode
    assert body["applicationMaxAgeHours"] == "1"
    assert body["priceBook"] == {
        "source": "demo-extract",
        "snapshotId": pricing_engine.manifest.snapshot_id,
        "pricedAsOf": pricing_engine.manifest.priced_as_of,
        "nonProduction": pricing_engine.manifest.non_production,
        "stale": False,
        "staleAfterDays": 30,
    }
    assert body["productLabel"] == "public-list run-rate benchmark"


@pytest.mark.parametrize(
    ("priced_as_of", "non_production", "expected"),
    [
        ("2026-08-29", False, False),  # exactly 30 days old
        ("2026-08-28", False, True),  # 31 days old
        ("2026-08-28T10:00:00Z", False, True),
        ("2020-01-01", True, False),  # frozen demo sets never go stale
        ("not a date", False, True),  # unreadable Published date shows the warning
    ],
)
def test_stale_harvest_rule_is_server_side(
    priced_as_of: str,
    non_production: bool,
    expected: bool,
) -> None:
    from datetime import date

    from app.routes import is_stale_harvest

    assert is_stale_harvest(priced_as_of, non_production, date(2026, 9, 28)) is expected


def test_capabilities_fail_closed_on_invalid_presentation_flag(
    client: TestClient,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setenv("PRESENT_AWS_PRICING", "yes")
    assert client.get("/api/capabilities").status_code == 503


# ---------------------------------------------------------------- verdict confidence


def test_verdict_confidence_draft_then_placeholder(
    client: TestClient,
    repository_root: Path,
) -> None:
    detail = _upload(client, _sample(repository_root))
    draft = detail["comparison"]
    assert draft["state"] == "DraftBenchmark"
    assert draft["verdict_confidence"] == "Draft"
    assert draft["assumption_sensitivities"] == []
    assert draft["headline_delta_percent"] is None

    ready = _resolve_all(client, detail)["comparison"]
    assert ready["state"] == "ReviewBaseline"
    assert ready["verdict_confidence"] == "Placeholder"


def test_verdict_confidence_final_without_placeholder_prices(
    client: TestClient,
    repository_root: Path,
    tmp_path: Path,
) -> None:
    path = _variant(repository_root, tmp_path, UBUNTU_EDITS, "ubuntu.xlsx")
    comparison = _resolve_all(client, _upload(client, path))["comparison"]
    assert comparison["state"] == "ReviewBaseline"
    assert not any(line["demo_assumption"] for line in comparison["line_items"])
    assert comparison["verdict_confidence"] == "Final"
    assert comparison["assumption_sensitivities"] == []


# ---------------------------------------------------------------- stress test


def test_assumption_stress_test_on_sample(
    client: TestClient,
    repository_root: Path,
) -> None:
    comparison = _resolve_all(client, _upload(client, _sample(repository_root)))["comparison"]
    assert comparison["aws_monthly_total"] == "3113.73"
    assert comparison["azure_monthly_total"] == "3135.15"
    assert comparison["cheaper_cloud"] == "AWS"
    assert comparison["headline_delta"] == "21.42"
    # 21.42 / 3135.15 * 100, the same figure as the List parity breakeven point.
    assert comparison["headline_delta_percent"] == "0.68"
    list_parity = next(
        item for item in comparison["breakeven_sensitivities"] if item["offer"] == "List"
    )
    assert (
        list_parity["points"][0]["target_discount_to_parity_percent"]
        == comparison["headline_delta_percent"]
    )

    [sensitivity] = comparison["assumption_sensitivities"]
    placeholder_lines = [
        line for line in comparison["line_items"] if line["demo_assumption"]
    ]
    assert sensitivity["line_ids"] == [line["id"] for line in placeholder_lines]
    assert sensitivity["line_ids"] == ["compute-2.rhel", "compute-3.rhel"]
    assert sensitivity["provider"] == "Azure"
    assert sensitivity["label"] == "Red Hat Enterprise Linux licensing"
    assert sensitivity["commercial_view"] == "List"
    assert sensitivity["assumed_monthly_amount"] == "200.96"
    assert sum(
        Decimal(line["azure_amount"]) for line in placeholder_lines
    ) == Decimal("200.96")

    assert sensitivity["at_zero"] == {
        "multiplier": "0",
        "assumed_monthly_amount": "0.00",
        "aws_monthly_total": "3113.73",
        "azure_monthly_total": "2934.19",
        "cheaper_cloud": "Azure",
        "headline_delta": "179.54",
        "changes_cheaper_cloud": True,
    }
    assert sensitivity["at_double"] == {
        "multiplier": "2",
        "assumed_monthly_amount": "401.92",
        "aws_monthly_total": "3113.73",
        "azure_monthly_total": "3336.11",
        "cheaper_cloud": "AWS",
        "headline_delta": "222.38",
        "changes_cheaper_cloud": False,
    }
    assert sensitivity["changes_cheaper_cloud"] is True


def _resolved_repository(repository_root: Path) -> tuple[InMemoryApplicationRepository, object]:
    repository = InMemoryApplicationRepository()
    detail = repository.create_from_intake(
        "fixture.xlsx",
        _sample(repository_root).read_bytes(),
    )
    repository.resolve_gap(
        detail.application.id,
        "P1.regions",
        GapResolutionRequest(
            resolved_by="Test Reviewer",
            azure_region="eastus2",
            aws_region="us-east-1",
        ),
    )
    repository.resolve_gap(
        detail.application.id,
        "compute-3.runtime",
        GapResolutionRequest(resolved_by="Test Reviewer", runtime_hours_month=Decimal("220")),
    )
    return repository, repository._records[detail.application.id]


def _recompare(record: object):
    return pricing_engine.compare(
        record.application,
        record.normalized.compute_units,
        record.normalized.database_units,
        record.normalized.storage_units,
        record.normalized.gaps,
    )


def test_assumption_sensitivities_are_covered_by_run_hash(
    repository_root: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    _, record = _resolved_repository(repository_root)
    baseline = _recompare(record)
    assert _recompare(record).run_hash == baseline.run_hash

    monkeypatch.setitem(
        pricing.PLACEHOLDER_COMPONENTS,
        "RHEL subscription uplift",
        (pricing.Cloud.AZURE, "Relabelled placeholder"),
    )
    relabelled = _recompare(record)
    assert relabelled.assumption_sensitivities[0].label == "Relabelled placeholder"
    assert relabelled.run_hash != baseline.run_hash


def test_unmapped_placeholder_component_fails_closed(
    repository_root: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    _, record = _resolved_repository(repository_root)
    monkeypatch.delitem(pricing.PLACEHOLDER_COMPONENTS, "RHEL subscription uplift")
    with pytest.raises(PricingConfigurationError, match="no stress-test mapping"):
        _recompare(record)


# ---------------------------------------------------------------- list summary


def test_application_list_returns_summary_fields(
    client: TestClient,
    repository_root: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.delenv("APPLICATION_MAX_AGE_HOURS", raising=False)
    draft = _upload(client, _sample(repository_root))
    ready = _resolve_all(client, _upload(client, _sample(repository_root)))

    listed = {item["id"]: item for item in client.get("/api/applications").json()}
    draft_row = listed[draft["application"]["id"]]
    assert draft_row["comparison_state"] == "DraftBenchmark"
    assert draft_row["verdict_confidence"] == "Draft"
    assert draft_row["open_question_count"] == 2
    assert draft_row["first_open_prompt"] == draft["gaps"][0]["prompt"]
    assert draft_row["cheaper_cloud"] is None
    assert draft_row["headline_delta"] is None
    assert draft_row["aws_monthly_total"] is None
    assert draft_row["azure_monthly_total"] is None
    assert draft_row["priced_as_of"] == pricing_engine.manifest.priced_as_of
    created = datetime.fromisoformat(draft_row["created_at"])
    assert datetime.fromisoformat(draft_row["expires_at"]) == created + timedelta(hours=1)

    ready_row = listed[ready["application"]["id"]]
    comparison = ready["comparison"]
    assert ready_row["comparison_state"] == "ReviewBaseline"
    assert ready_row["open_question_count"] == 0
    assert ready_row["first_open_prompt"] is None
    assert ready_row["verdict_confidence"] == comparison["verdict_confidence"] == "Placeholder"
    for field in (
        "cheaper_cloud",
        "headline_delta",
        "headline_delta_percent",
        "aws_monthly_total",
        "azure_monthly_total",
        "priced_as_of",
    ):
        assert ready_row[field] == comparison[field], field
    assert ready_row["excluded_cost_categories"] == [
        item["category"] for item in comparison["excluded_cost_ledger"]
    ]
    assert ready_row["excluded_cost_categories"]
    assert ready_row["placeholder_providers"] == ["Azure"]
    assert draft_row["placeholder_providers"] == []
    assert draft_row["excluded_cost_categories"]
    # The original Application fields are still present.
    for field in ("id", "name", "intake_file_name", "comparison_state", "created_at"):
        assert field in ready_row


def test_detail_reports_intake_summary_counts(
    client: TestClient,
    repository_root: Path,
) -> None:
    detail = _upload(client, _sample(repository_root))
    assert detail["intake_summary"] == {
        "server_count": 5,
        "server_vcpu": 28,
        "database_count": 1,
        "storage_count": 1,
        "storage_allocated_gb": "2000",
        "open_question_count": 2,
    }


# ---------------------------------------------------------------- workbook view


@pytest.mark.parametrize("flag", ["true", "false"])
def test_workbook_view_selects_layout_independent_of_default(
    client: TestClient,
    repository_root: Path,
    monkeypatch: pytest.MonkeyPatch,
    flag: str,
) -> None:
    monkeypatch.setenv("PRESENT_AWS_PRICING", flag)
    ready = _resolve_all(client, _upload(client, _sample(repository_root)))
    application_id = ready["application"]["id"]
    url = f"/api/applications/{application_id}/comparison/workbook"

    def summary_text(view: str | None) -> set[str]:
        response = client.get(url, params={"view": view} if view else None)
        assert response.status_code == 200, response.text
        sheet = load_workbook(io.BytesIO(response.content))["Pricing Summary"]
        return {
            cell.value
            for row in sheet.iter_rows()
            for cell in row
            if isinstance(cell.value, str)
        }

    both = summary_text("both")
    azure = summary_text("azure")
    aws = summary_text("aws")
    assert "Target discount to parity" in both
    assert "Azure discount to parity" in azure
    assert "Target discount to parity" not in azure
    assert not any("AWS" in value for value in azure)
    assert "AWS discount to parity" in aws
    assert not any("Azure" in value for value in aws)
    assert summary_text(None) == (both if flag == "true" else azure)


def test_aws_workbook_mirrors_azure_layout_with_aws_amounts(
    client: TestClient,
    repository_root: Path,
) -> None:
    ready = _resolve_all(client, _upload(client, _sample(repository_root)))
    comparison = ready["comparison"]
    url = f"/api/applications/{ready['application']['id']}/comparison/workbook"

    def sheets(view: str):
        response = client.get(url, params={"view": view})
        assert response.status_code == 200, response.text
        return load_workbook(io.BytesIO(response.content))

    aws_book = sheets("aws")
    azure_book = sheets("azure")
    scenarios = {
        (item["provider"], item["offer"]): item for item in comparison["commercial_scenarios"]
    }
    aws_summary = aws_book["Pricing Summary"]
    azure_summary = azure_book["Pricing Summary"]
    assert aws_summary["A1"].value == "AWS Pricing Summary"
    # Same cells, same formulas, only the provider's amounts differ.
    for row, offer in ((9, "List"), (10, "SavingsPlan"), (11, "Reservation")):
        assert aws_summary.cell(row=row, column=2).value == azure_summary.cell(row=row, column=2).value
        assert aws_summary.cell(row=row, column=4).value == azure_summary.cell(row=row, column=4).value
        expected = scenarios[("AWS", offer)]["monthly_total"]
        assert Decimal(str(aws_summary.cell(row=row, column=3).value)) == Decimal(expected)

    detail = aws_book["Pricing Detail"]
    assert detail["E5"].value == "AWS quantity"
    lines = comparison["line_items"]
    for row, line in enumerate(lines, start=6):
        rate = line["aws_commercial"]["list_rate"]
        cell = detail.cell(row=row, column=7).value
        assert (cell is None) if rate is None else Decimal(str(cell)) == Decimal(rate)
    detail_text = [
        cell.value
        for row in detail.iter_rows()
        for cell in row
        if isinstance(cell.value, str)
    ]
    # The only placeholder in the sample is on Azure; the AWS layout must not flag it.
    assert not any("demo assumption" in value for value in detail_text)
    assert not any(
        cell.fill.fgColor.rgb in {"FFFFF2CC", "00FFF2CC"}
        for row in detail.iter_rows(min_row=6)
        for cell in row
    )

    assumptions = [
        cell.value
        for row in aws_book["Pricing Assumptions"].iter_rows()
        for cell in row
        if isinstance(cell.value, str)
    ]
    assert "AWS region" in assumptions
    assert "Azure region" not in assumptions
    assert "AWS treatment" in assumptions
    evidence = aws_book["Pricing Evidence"]
    providers = {
        str(row[2].value).casefold()
        for row in evidence.iter_rows(min_row=1)
        if str(row[2].value).casefold() in {"aws", "azure"}
    }
    assert providers == {"aws"}
    assert aws_book["Pricing Exclusions"].max_row == azure_book["Pricing Exclusions"].max_row


def test_workbook_view_rejects_unknown_value(
    client: TestClient,
    repository_root: Path,
) -> None:
    ready = _resolve_all(client, _upload(client, _sample(repository_root)))
    response = client.get(
        f"/api/applications/{ready['application']['id']}/comparison/workbook",
        params={"view": "gcp"},
    )
    assert response.status_code == 422


# ---------------------------------------------------------------- evidence


def _run_hash_from_evidence(bundle: dict) -> str:
    facts = bundle["normalized_facts"]
    comparison = bundle["comparison"]
    application = {
        key: value
        for key, value in bundle["application"].items()
        if key != "comparison_state"
    }
    payload = {
        "application": application,
        "compute_units": facts["compute_units"],
        "database_units": facts["database_units"],
        "storage_units": facts["storage_units"],
        "gaps": facts["gaps"],
        "assumptions": comparison["assumptions"],
        "skumap_digest": comparison["skumap_content_digest"],
        "skumap_approval": comparison["skumap_approval"],
        "pricebook_hash": comparison["pricebook_content_hash"],
        "source_pricebook_hash": comparison["source_pricebook_content_hash"],
        "rule_version": bundle["calculator_rule_version"],
        "lines": comparison["line_items"],
        "commercial_scenarios": comparison["commercial_scenarios"],
        "breakeven_sensitivities": comparison["breakeven_sensitivities"],
        "headline_delta_percent": comparison["headline_delta_percent"],
        "verdict_confidence": comparison["verdict_confidence"],
        "assumption_sensitivities": comparison["assumption_sensitivities"],
        "outputs": {
            "state": comparison["state"],
            "can_export": comparison["can_export"],
            "aws_monthly_total": comparison["aws_monthly_total"],
            "azure_monthly_total": comparison["azure_monthly_total"],
            "headline_delta": comparison["headline_delta"],
            "cheaper_cloud": comparison["cheaper_cloud"],
            "cost_drivers": comparison["cost_drivers"],
            "excluded_cost_ledger": comparison["excluded_cost_ledger"],
        },
    }
    return pricing._digest(payload)


def test_evidence_json_alone_reproduces_the_run_hash(
    client: TestClient,
    repository_root: Path,
) -> None:
    ready = _resolve_all(client, _upload(client, _sample(repository_root)))
    response = client.get(
        f"/api/applications/{ready['application']['id']}/comparison/export"
    )
    assert response.status_code == 200
    bundle = json.loads(response.content)
    assert set(bundle) == {
        "schema_version",
        "calculator_rule_version",
        "application",
        "normalized_facts",
        "comparison",
    }
    assert set(bundle["comparison"]) == set(ready["comparison"])
    assert bundle["comparison"]["verdict_confidence"] == "Placeholder"
    assert bundle["comparison"]["assumption_sensitivities"]
    assert bundle["comparison"]["headline_delta_percent"] is not None
    assert "expires_at" not in bundle["application"]
    assert _run_hash_from_evidence(bundle) == bundle["comparison"]["run_hash"]
    assert bundle["comparison"]["run_hash"] == ready["comparison"]["run_hash"]


@pytest.mark.parametrize(
    ("field", "value"),
    [
        ("verdict_confidence", "Final"),
        ("cheaper_cloud", "Azure"),
        ("aws_monthly_total", "1.00"),
        ("azure_monthly_total", "1.00"),
        ("headline_delta", "0.00"),
        ("state", "DraftBenchmark"),
        ("can_export", False),
        ("excluded_cost_ledger", []),
        ("cost_drivers", []),
    ],
)
def test_evidence_run_hash_breaks_when_a_field_is_altered(
    client: TestClient,
    repository_root: Path,
    field: str,
    value: object,
) -> None:
    ready = _resolve_all(client, _upload(client, _sample(repository_root)))
    bundle = json.loads(
        client.get(
            f"/api/applications/{ready['application']['id']}/comparison/export"
        ).content
    )
    assert bundle["comparison"][field] != value
    bundle["comparison"][field] = value
    assert _run_hash_from_evidence(bundle) != bundle["comparison"]["run_hash"]
