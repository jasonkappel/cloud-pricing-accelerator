import io
import json
import re
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import load_workbook

from app.intake import IntakeValidationError, parse_intake, validate_openxml_package


def test_fixture_normalizes_to_expected_contract(repository_root: Path) -> None:
    content = (repository_root / "samples" / "synthetic_intake_completed.xlsx").read_bytes()
    expected = json.loads(
        (repository_root / "samples" / "synthetic_intake_completed_expected.json").read_text(
            encoding="utf-8"
        )
    )

    normalized = parse_intake(content)

    assert normalized.application_name == expected["application"]
    assert normalized.environments == expected["environments"]
    assert [
        {
            "name": unit.name,
            "env": unit.environment,
            "count": unit.count,
            "os": unit.operating_system,
            "license": unit.license_model,
            "vcpu": unit.vcpu_each,
            "ram_gb": int(unit.ram_gb_each),
            "runtime": (
                "24x7" if unit.runtime_hours_month is not None else "Unknown"
            ),
        }
        for unit in normalized.compute_units
    ] == expected["compute_units"]
    assert [gap.id for gap in normalized.gaps] == [
        gap["id"] for gap in expected["expected_gaps"]
    ]


def test_ole2_sample_is_rejected_with_resave_message(repository_root: Path) -> None:
    content = (repository_root / "samples" / "reject_ole2_sample.bin").read_bytes()

    with pytest.raises(IntakeValidationError, match="Save As .xlsx"):
        validate_openxml_package(content)


def test_non_workbook_zip_is_rejected() -> None:
    stream = io.BytesIO()
    with ZipFile(stream, "w", ZIP_DEFLATED) as archive:
        archive.writestr("hello.txt", "not a workbook")

    with pytest.raises(IntakeValidationError, match="required workbook parts"):
        validate_openxml_package(stream.getvalue())


def test_high_compression_ratio_is_rejected() -> None:
    stream = io.BytesIO()
    with ZipFile(stream, "w", ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "<Types />")
        archive.writestr("_rels/.rels", "<Relationships />")
        archive.writestr("xl/workbook.xml", "<workbook />")
        archive.writestr("xl/_rels/workbook.xml.rels", "<Relationships />")
        archive.writestr("xl/worksheets/sheet1.xml", "0" * 500_000)

    with pytest.raises(IntakeValidationError, match="compression ratio"):
        validate_openxml_package(stream.getvalue())


def test_external_link_part_is_rejected(repository_root: Path) -> None:
    source = repository_root / "samples" / "synthetic_intake_completed.xlsx"
    stream = io.BytesIO()
    with ZipFile(source) as original, ZipFile(stream, "w", ZIP_DEFLATED) as modified:
        for member in original.infolist():
            modified.writestr(member, original.read(member.filename))
        modified.writestr("xl/externalLinks/externalLink1.xml", "<externalLink />")

    with pytest.raises(IntakeValidationError, match="external workbook links"):
        validate_openxml_package(stream.getvalue())


def test_macro_part_check_is_case_insensitive(repository_root: Path) -> None:
    source = repository_root / "samples" / "synthetic_intake_completed.xlsx"
    stream = io.BytesIO()
    with ZipFile(source) as original, ZipFile(stream, "w", ZIP_DEFLATED) as modified:
        for member in original.infolist():
            modified.writestr(member, original.read(member.filename))
        modified.writestr("XL/VBAPROJECT.BIN", b"macro")

    with pytest.raises(IntakeValidationError, match="Macros"):
        validate_openxml_package(stream.getvalue())


def test_dtd_or_entity_declarations_are_rejected(repository_root: Path) -> None:
    source = repository_root / "samples" / "synthetic_intake_completed.xlsx"
    stream = io.BytesIO()
    with ZipFile(source) as original, ZipFile(stream, "w", ZIP_DEFLATED) as modified:
        for member in original.infolist():
            content = original.read(member.filename)
            if member.filename == "xl/workbook.xml":
                content = content.replace(
                    b"<workbook ",
                    b'<!DOCTYPE workbook [<!ENTITY x "boom">]><workbook ',
                    1,
                )
            modified.writestr(member, content)

    with pytest.raises(IntakeValidationError, match="DTD or entity"):
        validate_openxml_package(stream.getvalue())


def test_external_relationship_target_is_rejected(repository_root: Path) -> None:
    source = repository_root / "samples" / "synthetic_intake_completed.xlsx"
    stream = io.BytesIO()
    relationship = (
        b'<Relationship Id="external" '
        b'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" '
        b'Target="https://example.invalid/workbook.xlsx" TargetMode="External"/>'
    )
    with ZipFile(source) as original, ZipFile(stream, "w", ZIP_DEFLATED) as modified:
        for member in original.infolist():
            content = original.read(member.filename)
            if member.filename == "xl/_rels/workbook.xml.rels":
                content = content.replace(b"</Relationships>", relationship + b"</Relationships>")
            modified.writestr(member, content)

    with pytest.raises(IntakeValidationError, match="External workbook relationships"):
        validate_openxml_package(stream.getvalue())


def test_corrupt_compressed_member_is_a_typed_rejection(
    repository_root: Path,
) -> None:
    content = bytearray(
        (repository_root / "samples" / "synthetic_intake_completed.xlsx").read_bytes()
    )
    with ZipFile(io.BytesIO(content)) as archive:
        member = archive.getinfo("xl/worksheets/sheet2.xml")
        header = member.header_offset
        filename_length = int.from_bytes(content[header + 26 : header + 28], "little")
        extra_length = int.from_bytes(content[header + 28 : header + 30], "little")
        data_start = header + 30 + filename_length + extra_length
        content[data_start + member.compress_size // 2] ^= 0xFF

    with pytest.raises(IntakeValidationError, match="valid OpenXML package"):
        validate_openxml_package(bytes(content))


def test_oversized_worksheet_dimensions_are_rejected(
    repository_root: Path,
) -> None:
    source = repository_root / "samples" / "synthetic_intake_completed.xlsx"
    stream = io.BytesIO()
    with ZipFile(source) as original, ZipFile(stream, "w", ZIP_DEFLATED) as modified:
        for member in original.infolist():
            content = original.read(member.filename)
            if member.filename == "xl/worksheets/sheet2.xml":
                content = re.sub(
                    rb'<dimension ref="[^"]+"',
                    b'<dimension ref="A1:XFD1048576"',
                    content,
                    count=1,
                )
            modified.writestr(member, content)

    with pytest.raises(IntakeValidationError, match="dimension limits"):
        parse_intake(stream.getvalue())


def test_drawing_parts_are_rejected(repository_root: Path) -> None:
    source = repository_root / "samples" / "synthetic_intake_completed.xlsx"
    stream = io.BytesIO()
    with ZipFile(source) as original, ZipFile(stream, "w", ZIP_DEFLATED) as modified:
        for member in original.infolist():
            modified.writestr(member, original.read(member.filename))
        modified.writestr("xl/drawings/drawing1.xml", "<xdr:wsDr />")

    with pytest.raises(IntakeValidationError, match="drawings"):
        parse_intake(stream.getvalue())


def test_custom_xml_parts_are_rejected(repository_root: Path) -> None:
    source = repository_root / "samples" / "synthetic_intake_completed.xlsx"
    stream = io.BytesIO()
    with ZipFile(source) as original, ZipFile(stream, "w", ZIP_DEFLATED) as modified:
        for member in original.infolist():
            modified.writestr(member, original.read(member.filename))
        modified.writestr("customXml/item1.xml", "<root />")

    with pytest.raises(IntakeValidationError, match="unsupported OpenXML parts"):
        parse_intake(stream.getvalue())


def test_missing_worksheet_dimension_is_rejected(repository_root: Path) -> None:
    source = repository_root / "samples" / "synthetic_intake_completed.xlsx"
    stream = io.BytesIO()
    with ZipFile(source) as original, ZipFile(stream, "w", ZIP_DEFLATED) as modified:
        for member in original.infolist():
            content = original.read(member.filename)
            if member.filename == "xl/worksheets/sheet2.xml":
                content = re.sub(rb"<dimension[^>]*/>", b"", content, count=1)
            modified.writestr(member, content)

    with pytest.raises(IntakeValidationError, match="bounded dimension"):
        parse_intake(stream.getvalue())


def test_underdeclared_worksheet_dimension_is_rejected(
    repository_root: Path,
) -> None:
    source = repository_root / "samples" / "synthetic_intake_completed.xlsx"
    stream = io.BytesIO()
    with ZipFile(source) as original, ZipFile(stream, "w", ZIP_DEFLATED) as modified:
        for member in original.infolist():
            content = original.read(member.filename)
            if member.filename == "xl/worksheets/sheet2.xml":
                content = re.sub(
                    rb'<dimension ref="[^"]+"',
                    b'<dimension ref="A1:K2"',
                    content,
                    count=1,
                )
            modified.writestr(member, content)

    with pytest.raises(IntakeValidationError, match="outside its declared dimension"):
        parse_intake(stream.getvalue())


def test_fractional_integer_fields_are_rejected(
    repository_root: Path,
    tmp_path: Path,
) -> None:
    workbook = load_workbook(
        repository_root / "samples" / "synthetic_intake_completed.xlsx"
    )
    workbook["2 Servers"]["C2"] = 1.5
    path = tmp_path / "fractional-count.xlsx"
    workbook.save(path)

    with pytest.raises(IntakeValidationError, match="whole number"):
        parse_intake(path.read_bytes())


def test_partial_rows_fail_as_validation_errors(
    repository_root: Path,
    tmp_path: Path,
) -> None:
    workbook = load_workbook(
        repository_root / "samples" / "synthetic_intake_completed.xlsx"
    )
    for cell in workbook["2 Servers"][4]:
        cell.value = None
    workbook["2 Servers"]["A4"] = "partial-server"
    path = tmp_path / "partial-row.xlsx"
    workbook.save(path)

    with pytest.raises(IntakeValidationError, match="missing"):
        parse_intake(path.read_bytes())


def test_supported_p1_answer_is_canonicalized(
    repository_root: Path,
    tmp_path: Path,
) -> None:
    workbook = load_workbook(
        repository_root / "samples" / "synthetic_intake_completed.xlsx"
    )
    platform = workbook["Platform + FinOps"]
    for row in platform.iter_rows(min_row=2):
        if row[0].value == "P1":
            row[4].value = "Azure eastus2 and AWS us-east-1"
    path = tmp_path / "regions.xlsx"
    workbook.save(path)

    normalized = parse_intake(path.read_bytes())
    region_gap = next(gap for gap in normalized.gaps if gap.id == "P1.regions")
    assert region_gap.status.value == "Resolved"
    assert region_gap.canonical_value == {
        "azure_region": "eastus2",
        "aws_region": "us-east-1",
    }


def test_region_identifiers_require_exact_token_boundaries(
    repository_root: Path,
    tmp_path: Path,
) -> None:
    workbook = load_workbook(
        repository_root / "samples" / "synthetic_intake_completed.xlsx"
    )
    platform = workbook["Platform + FinOps"]
    for row in platform.iter_rows(min_row=2):
        if row[0].value == "P1":
            row[4].value = "Azure eastus2euap and AWS us-east-1"
    path = tmp_path / "near-match-regions.xlsx"
    workbook.save(path)

    normalized = parse_intake(path.read_bytes())
    region_gap = next(gap for gap in normalized.gaps if gap.id == "P1.regions")
    assert region_gap.status.value == "Open"


def test_duplicate_names_get_distinct_gap_ids(
    repository_root: Path,
    tmp_path: Path,
) -> None:
    workbook = load_workbook(
        repository_root / "samples" / "synthetic_intake_completed.xlsx"
    )
    servers = workbook["2 Servers"]
    servers["A2"] = "duplicate"
    servers["A3"] = "duplicate"
    servers["I2"] = "Unknown"
    servers["I3"] = "Unknown"
    path = tmp_path / "duplicate-names.xlsx"
    workbook.save(path)

    normalized = parse_intake(path.read_bytes())
    runtime_gap_ids = [
        gap.id for gap in normalized.gaps if gap.kind.value == "RuntimeHours"
    ]
    assert runtime_gap_ids == [
        "compute-1.runtime",
        "compute-2.runtime",
        "compute-3.runtime",
    ]


def test_ha_node_count_contradictions_are_rejected(
    repository_root: Path,
    tmp_path: Path,
) -> None:
    workbook = load_workbook(
        repository_root / "samples" / "synthetic_intake_completed.xlsx"
    )
    workbook["3 Databases"]["E2"] = 1
    path = tmp_path / "ha-contradiction.xlsx"
    workbook.save(path)

    with pytest.raises(IntakeValidationError, match="at least two nodes"):
        parse_intake(path.read_bytes())


def test_untrusted_text_is_neutralized_consistently(
    repository_root: Path,
    tmp_path: Path,
) -> None:
    workbook = load_workbook(
        repository_root / "samples" / "synthetic_intake_completed.xlsx"
    )
    workbook["2 Servers"]["A4"] = "+spreadsheet-command"
    path = tmp_path / "untrusted-text.xlsx"
    workbook.save(path)

    normalized = parse_intake(path.read_bytes())
    assert normalized.compute_units[2].name == "'+spreadsheet-command"
    assert "'+spreadsheet-command" in normalized.gaps[0].prompt


def test_empty_workload_is_rejected(
    repository_root: Path,
    tmp_path: Path,
) -> None:
    workbook = load_workbook(
        repository_root / "samples" / "synthetic_intake_completed.xlsx"
    )
    for sheet_name in ("2 Servers", "3 Databases", "4 Storage"):
        sheet = workbook[sheet_name]
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.value = None
    platform = workbook["Platform + FinOps"]
    for row in platform.iter_rows(min_row=2):
        if row[0].value == "P1":
            row[4].value = "Azure eastus2 and AWS us-east-1"
    path = tmp_path / "empty-workload.xlsx"
    workbook.save(path)

    with pytest.raises(IntakeValidationError, match="at least one"):
        parse_intake(path.read_bytes())
