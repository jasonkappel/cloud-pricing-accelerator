"""
generate_fixtures.py  —  build clean, fully synthetic test fixtures for the ingestion path.

Run: py generate_fixtures.py

Produces (all synthetic, no customer data, no MIP labels, no personal metadata):
  - synthetic_intake_completed.xlsx      a completed, in-scope intake (VM + PostgreSQL + block)
  - synthetic_intake_completed_expected.json   expected normalized facts + expected Gaps
  - reject_ole2_sample.bin               a minimal file carrying the OLE2 signature (reject-path test)

Design intent (see the README "Design rules"):
  - Everything is in v1 MVP scope: IaaS VMs (Windows + Linux), managed PostgreSQL, attached block storage.
  - Two deliberate open Gaps so the build can exercise the CompletenessGate DraftBenchmark path:
      * AppFactGap: one server's Runtime is left "Unknown".
      * PricingPolicyGap: Platform/FinOps P1 (approved regions) is left open / Not started.
  - docProps are written clean by openpyxl (no author, no sensitivity label, no tenant/site id, no paths).
"""

from __future__ import annotations
import json
import openpyxl
from openpyxl.worksheet.table import Table, TableStyleInfo

APP = "StatementHub (synthetic)"

def build_workbook(path: str) -> None:
    wb = openpyxl.Workbook()

    # 1 App Basics (fixed-range section, NOT an Excel table)
    ws = wb.active
    ws.title = "1 App Basics"
    ws.append(["ID", "Question", "Your answer", "Example answer", "Who can help"])
    basics = [
        ("A1", "What is the application name?", APP, "Customer Statements", "Application owner"),
        ("A2", "In one sentence, what does the application do?",
         "Generates monthly account statements and serves them to customers.",
         "Generates customer statements and stores the PDFs.", "Application owner"),
        ("A3", "Which environments are needed?", "Dev, Test, Pre-Prod, and Production.",
         "Dev, Test, Pre-Prod, and Production.", "Application owner"),
        ("A4", "Who should we contact with questions?",
         "App Support team; Infrastructure team; DBA team.",
         "Jane Smith, App Support; John Doe, Infrastructure.", "Application owner"),
        ("A5", "What recovery is required?", "99.9% availability; RPO 15 minutes; RTO 2 hours.",
         "99.9% availability; RPO 15 minutes; RTO 2 hours.", "Application owner / resilience"),
        ("A6", "Are there important vendor, software-license, compliance, or technology restrictions?",
         "Data must remain in approved US regions.", "Oracle BYOL; vendor requires x86.",
         "Technical / software asset contact"),
        ("A7", "Does the application connect to any third parties? If yes, who and how?",
         "No third-party connections.", "Yes, payment provider over private connectivity.",
         "Application / network contact"),
    ]
    for row in basics:
        ws.append(list(row))

    # 2 Servers (Excel table). In-scope IaaS VMs, Windows + Linux. One Runtime left Unknown (AppFactGap).
    ws = wb.create_sheet("2 Servers")
    servers_header = ["Server / group name", "Environment", "Count", "Current platform",
                      "OS / edition", "License model", "vCPU each", "RAM GB each", "Runtime",
                      "Monitoring report / notes (optional)", "Source and date"]
    ws.append(servers_header)
    servers = [
        ["web-prod", "Production", 2, "VMware", "Windows Server 2022", "License included", 4, 16,
         "24x7", "vCenter: CPU p95 40%; memory p95 65%", "CMDB + vCenter, 2026-09-01"],
        ["app-prod", "Production", 2, "VMware", "RHEL 9", "Subscription", 8, 32,
         "24x7", "vCenter: CPU p95 55%; memory p95 70%", "CMDB + vCenter, 2026-09-01"],
        ["app-test", "Test", 1, "VMware", "RHEL 9", "Subscription", 4, 16,
         "Unknown", "Runtime schedule not confirmed", "CMDB, 2026-09-01"],
    ]
    for r in servers:
        ws.append(r)
    _add_table(ws, "T2Servers", len(servers_header), len(servers) + 1)

    # 3 Databases (Excel table). PostgreSQL only (in scope).
    ws = wb.create_sheet("3 Databases")
    db_header = ["Database name", "Environment", "Engine / version", "License model",
                 "Instances / nodes", "vCPU & RAM each", "Database size GB",
                 "HA / backup / recovery", "Monitoring report / notes (optional)", "Source and date"]
    ws.append(db_header)
    dbs = [
        ["STMTPG", "Production", "PostgreSQL 15", "Open source / no fee", 2, "8 vCPU / 64 GB", 1000,
         "HA replica; daily backup; RPO 15 min; RTO 2 hr", "pg_stat report 2026-09-01",
         "DBA inventory, 2026-09-01"],
    ]
    for r in dbs:
        ws.append(r)
    _add_table(ws, "T3Databases", len(db_header), len(dbs) + 1)

    # 4 Storage (Excel table). Attached block storage, with capacity + IOPS in notes.
    ws = wb.create_sheet("4 Storage")
    st_header = ["Storage name", "Environment", "Type", "Protocol", "Allocated GB", "Used GB",
                 "Backup / snapshot / retention", "Performance report / notes (optional)",
                 "Source and date"]
    ws.append(st_header)
    storage = [
        ["statements-prod-disk", "Production", "Server disk / block", "Block", 2000, 1400,
         "Daily snapshot; 35-day retention", "Storage report: peak 3000 IOPS, 180 MB/s",
         "Disk report, 2026-09-01"],
    ]
    for r in storage:
        ws.append(r)
    _add_table(ws, "T4Storage", len(st_header), len(storage) + 1)

    # Platform + FinOps (fixed-range section). P1 left OPEN -> PricingPolicyGap.
    ws = wb.create_sheet("Platform + FinOps")
    ws.append(["ID", "Owner", "Question", "Example / guidance", "Answer", "Source and date", "Status"])
    pf = [
        ("P1", "Platform", "Which primary and DR regions are approved?",
         "Primary and DR regions per policy.", "", "", "Not started"),   # OPEN (material)
        ("P2", "Platform / Network", "How will approved third-party connections be implemented?",
         "PrivateLink, VPN, or no new component.", "No new component (A7 = none).", "2026-09-02", "Complete"),
        ("P3", "Platform / Security", "Which mandatory shared platform and security services apply?",
         "Logging, backup, DNS, keys, secrets.", "Standard logging, backup, DNS, secrets.",
         "2026-09-02", "Complete"),
        ("F1", "FinOps", "Confirm currency, estimate horizon, and pricing basis.",
         "USD, 36 months, On-Demand baseline.", "USD, 36 months, On-Demand baseline.",
         "2026-09-02", "Complete"),
        ("F4", "FinOps / SAM", "Confirm license treatment and BYOL eligibility.",
         "Windows license included; open-source no fee.", "Windows included; PostgreSQL open source.",
         "2026-09-02", "Complete"),
    ]
    for r in pf:
        ws.append(list(r))

    wb.save(path)


def _add_table(ws, name: str, ncols: int, nrows: int) -> None:
    last_col = openpyxl.utils.get_column_letter(ncols)
    ref = f"A1:{last_col}{nrows}"
    tbl = Table(displayName=name, ref=ref)
    tbl.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    ws.add_table(tbl)


def expected_output() -> dict:
    return {
        "application": APP,
        "environments": ["Dev", "Test", "Pre-Prod", "Production"],
        "compute_units": [
            {"name": "web-prod", "env": "Production", "count": 2, "os": "Windows Server 2022",
             "license": "License included", "vcpu": 4, "ram_gb": 16, "runtime": "24x7"},
            {"name": "app-prod", "env": "Production", "count": 2, "os": "RHEL 9",
             "license": "Subscription", "vcpu": 8, "ram_gb": 32, "runtime": "24x7"},
            {"name": "app-test", "env": "Test", "count": 1, "os": "RHEL 9",
             "license": "Subscription", "vcpu": 4, "ram_gb": 16, "runtime": "Unknown"},
        ],
        "database_units": [
            {"name": "STMTPG", "env": "Production", "engine": "PostgreSQL 15", "nodes": 2,
             "vcpu": 8, "ram_gb": 64, "size_gb": 1000, "ha": True},
        ],
        "storage_units": [
            {"name": "statements-prod-disk", "env": "Production", "type": "block",
             "allocated_gb": 2000, "used_gb": 1400, "target_iops": 3000, "target_mbps": 180},
        ],
        "expected_gaps": [
            {"id": "app-test.runtime", "type": "AppFactGap", "material": True,
             "reason": "Server runtime is Unknown; cannot compute run-hours."},
            {"id": "P1.regions", "type": "PricingPolicyGap", "material": True,
             "reason": "Approved primary/DR regions not set; cannot select a PriceBook region."},
        ],
        "expected_state": "DraftBenchmark",
        "note": "Two open material Gaps must hold the headline until resolved or defaulted by policy.",
    }


def build_ole2_reject(path: str) -> None:
    # Minimal file carrying the OLE2 (compound file) signature so the ingestion flow can test the
    # magic-byte reject path. This is a fixture, not a real workbook.
    sig = bytes([0xD0, 0xCF, 0x11, 0xE0, 0xA1, 0xB1, 0x1A, 0xE1])
    with open(path, "wb") as fh:
        fh.write(sig + b"\x00" * 504)


if __name__ == "__main__":
    import os
    here = os.path.dirname(os.path.abspath(__file__))
    xlsx = os.path.join(here, "synthetic_intake_completed.xlsx")
    js = os.path.join(here, "synthetic_intake_completed_expected.json")
    ole = os.path.join(here, "reject_ole2_sample.bin")
    build_workbook(xlsx)
    with open(js, "w", encoding="utf-8") as fh:
        json.dump(expected_output(), fh, indent=2)
    build_ole2_reject(ole)
    print("Wrote:", os.path.basename(xlsx), os.path.basename(js), os.path.basename(ole))
