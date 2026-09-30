import io
import posixpath
import re
import zlib
from datetime import UTC, datetime
from decimal import Decimal
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

from defusedxml import ElementTree as DefusedElementTree
from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_to_tuple, range_boundaries

from app.models import (
    BackupPolicy,
    ComputeUnit,
    DatabaseUnit,
    Gap,
    GapKind,
    GapStatus,
    NetworkTransferProfile,
    NormalizedIntake,
    ObservabilityProfile,
    SecurityProfile,
    StoragePerformanceProfile,
    StorageUnit,
    SupportTier,
    TopologyTier,
)

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_EXPANDED_BYTES = 20 * 1024 * 1024
MAX_MEMBER_BYTES = 10 * 1024 * 1024
MAX_MEMBER_COUNT = 256
MAX_COMPRESSION_RATIO = Decimal("100")
MAX_SHEET_ROWS = 10_000
MAX_SHEET_COLUMNS = 64
MAX_SHEET_CELLS = 100_000
OLE2_SIGNATURE = b"\xd0\xcf\x11\xe0"
OPENXML_SIGNATURES = (b"PK\x03\x04", b"PK\x05\x06", b"PK\x07\x08")
REQUIRED_PARTS = {
    "[Content_Types].xml",
    "xl/workbook.xml",
    "xl/_rels/workbook.xml.rels",
    "_rels/.rels",
}
FORBIDDEN_PART_PREFIXES = (
    "xl/externallinks/",
    "xl/charts/",
    "xl/drawings/",
    "xl/pivotcache/",
    "xl/pivottables/",
)
FORBIDDEN_PART_NAMES = {
    "encryptedpackage",
    "encryptioninfo",
    "xl/vbaproject.bin",
}
ALLOWED_PART_NAMES = {
    "[content_types].xml",
    "_rels/.rels",
    "docprops/app.xml",
    "docprops/core.xml",
    "xl/workbook.xml",
    "xl/_rels/workbook.xml.rels",
    "xl/styles.xml",
    "xl/sharedstrings.xml",
}
ALLOWED_PART_PREFIXES = (
    "xl/theme/",
    "xl/worksheets/",
    "xl/tables/",
)
FORMULA_PATTERN = re.compile(rb"<(?:[A-Za-z0-9_]+:)?f(?:\s|>)", re.IGNORECASE)
# Checked on the parsed XML of every part, so no encoding or part name can hide one.
# Cell, table, conditional-formatting, and data-validation formulas all recalculate in Excel.
FORMULA_ELEMENTS = {"f", "calculatedcolumnformula", "totalsrowformula", "formula", "formula1", "formula2"}
WORKSHEET_RELATIONSHIP = "/relationships/worksheet"
OFFICE_DOCUMENT_RELATIONSHIP = "/relationships/officedocument"
RELATIONSHIP_ID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
# Print areas, print titles, and filter ranges are the only defined names kept, and only as plain
# range references. Each piece is checked by a linear parser, never a nested regex.
SAFE_DEFINED_NAMES = {"_xlnm.print_area", "_xlnm.print_titles", "_xlnm._filterdatabase"}
MAX_DEFINED_NAME_LENGTH = 2048
UNQUOTED_SHEET = re.compile(r"[A-Za-z0-9_.]{1,31}")
CELL_RANGE = re.compile(
    r"\$?[A-Z]{1,3}\$?[0-9]{1,7}(?::\$?[A-Z]{1,3}\$?[0-9]{1,7})?"
    r"|\$?[A-Z]{1,3}:\$?[A-Z]{1,3}"
    r"|\$?[0-9]{1,7}:\$?[0-9]{1,7}"
)
WORKSHEET_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"
UNSAFE_XML_PATTERN = re.compile(rb"<!\s*(?:DOCTYPE|ENTITY)\b", re.IGNORECASE)
SHARED_EXCLUSION_POLICY = "PILOT-SHARED-PLATFORM-EXCLUSION-V1"
SUPPORT_EXCLUSION_POLICY = "PILOT-SUPPORT-EXCLUSION-V1"
OBSERVABILITY_EXCLUSION_POLICY = "PILOT-OBSERVABILITY-EXCLUSION-V1"
SECURITY_EXCLUSION_POLICY = "PILOT-SECURITY-EXCLUSION-V1"


class IntakeValidationError(ValueError):
    pass


def validate_openxml_package(content: bytes) -> None:
    if len(content) > MAX_UPLOAD_BYTES:
        raise IntakeValidationError("The Intake exceeds the 10 MiB upload limit.")
    if content.startswith(OLE2_SIGNATURE):
        raise IntakeValidationError(
            "Legacy binary format detected. Open in Excel and Save As .xlsx to continue."
        )
    if not content.startswith(OPENXML_SIGNATURES):
        raise IntakeValidationError("The Intake is not a valid OpenXML package.")

    try:
        with ZipFile(io.BytesIO(content)) as archive:
            members = archive.infolist()
            names = {member.filename for member in members}
            normalized_names = {name.casefold() for name in names}
            actual_names = {name.casefold(): name for name in names}
            if len(members) > MAX_MEMBER_COUNT:
                raise IntakeValidationError("The OpenXML package contains too many members.")
            if not {part.casefold() for part in REQUIRED_PARTS}.issubset(normalized_names):
                raise IntakeValidationError("The OpenXML package is missing required workbook parts.")
            if any(
                name.casefold() in FORBIDDEN_PART_NAMES
                or name.casefold().startswith(FORBIDDEN_PART_PREFIXES)
                or "vbaproject" in name.casefold()
                for name in names
            ):
                raise IntakeValidationError(
                    "Macros, encryption, external workbook links, drawings, charts, "
                    "and pivot content are not accepted."
                )
            unsupported_parts = [
                name
                for name in names
                if name.casefold() not in ALLOWED_PART_NAMES
                and not name.casefold().startswith(ALLOWED_PART_PREFIXES)
            ]
            if unsupported_parts:
                raise IntakeValidationError(
                    "The workbook contains unsupported OpenXML parts that cannot be "
                    "preserved safely."
                )

            expanded_size = 0
            for member in members:
                if member.file_size > MAX_MEMBER_BYTES:
                    raise IntakeValidationError("An OpenXML package member exceeds the size limit.")
                expanded_size += member.file_size
                if expanded_size > MAX_EXPANDED_BYTES:
                    raise IntakeValidationError("The expanded OpenXML package exceeds the size limit.")
                if member.file_size and member.compress_size == 0:
                    raise IntakeValidationError("The OpenXML package has an unsafe compression ratio.")
                if member.compress_size:
                    ratio = Decimal(member.file_size) / Decimal(member.compress_size)
                    if ratio > MAX_COMPRESSION_RATIO:
                        raise IntakeValidationError(
                            "The OpenXML package has an unsafe compression ratio."
                        )
                if member.is_dir():
                    continue
                if not member.filename.casefold().endswith((".xml", ".rels")):
                    raise IntakeValidationError(
                        "The workbook contains unsupported OpenXML parts that cannot be "
                        "preserved safely."
                    )
                xml_bytes = archive.read(member)
                if UNSAFE_XML_PATTERN.search(xml_bytes):
                    raise IntakeValidationError(
                        "OpenXML packages containing DTD or entity declarations are not accepted."
                    )
                root = _parse_part(xml_bytes)
                if FORMULA_PATTERN.search(xml_bytes) or any(
                    _local_name(element.tag) in FORMULA_ELEMENTS
                    or (
                        _local_name(element.tag) == "cfvo"
                        and str(element.get("type", "")).casefold() == "formula"
                    )
                    or str(element.get("totalsRowFunction", "none")).casefold() not in ("", "none")
                    for element in root.iter()
                ):
                    raise IntakeValidationError(
                        "Workbook formulas are not accepted in Intake fields."
                    )
                if (
                    member.filename.casefold().startswith("xl/worksheets/")
                    and member.filename.casefold().endswith(".xml")
                ):
                    _validate_worksheet_xml(xml_bytes)
                if member.filename.casefold().endswith(".rels"):
                    _validate_relationships(member.filename, xml_bytes)

            _validate_workbook_location(
                archive.read(actual_names["[content_types].xml"]), archive.read(actual_names["_rels/.rels"])
            )
            _validate_worksheet_targets(archive.read(actual_names["xl/_rels/workbook.xml.rels"]))
            _validate_workbook_sheets(
                archive.read(actual_names["xl/workbook.xml"]), archive.read(actual_names["xl/_rels/workbook.xml.rels"])
            )
            content_types = archive.read(actual_names["[content_types].xml"])
            folded_content_types = content_types.lower()
            if b"macroenabled" in folded_content_types or b"vbaproject" in folded_content_types:
                raise IntakeValidationError("Macro-enabled workbooks are not accepted.")
            for part in REQUIRED_PARTS:
                try:
                    DefusedElementTree.fromstring(
                        archive.read(actual_names[part.casefold()])
                    )
                except (ParseError, ValueError) as error:
                    raise IntakeValidationError(
                        "The OpenXML package contains invalid workbook relationships."
                    ) from error
    except IntakeValidationError:
        raise
    except (BadZipFile, zlib.error, RuntimeError, EOFError) as error:
        raise IntakeValidationError("The Intake is not a valid OpenXML package.") from error


def parse_intake(content: bytes) -> NormalizedIntake:
    validate_openxml_package(content)
    try:
        workbook = load_workbook(
            io.BytesIO(content),
            read_only=True,
            data_only=False,
            keep_links=False,
        )
    except Exception as error:
        raise IntakeValidationError("The OpenXML workbook could not be parsed safely.") from error

    required_sheets = {
        "1 App Basics",
        "2 Servers",
        "3 Databases",
        "4 Storage",
        "Platform + FinOps",
    }
    if not required_sheets.issubset(workbook.sheetnames):
        raise IntakeValidationError("The Intake does not match the required workbook structure.")
    for sheet in workbook.worksheets:
        if sheet.max_row is None or sheet.max_column is None:
            raise IntakeValidationError("A worksheet does not declare a bounded dimension.")
        if (
            sheet.max_row > MAX_SHEET_ROWS
            or sheet.max_column > MAX_SHEET_COLUMNS
            or sheet.max_row * sheet.max_column > MAX_SHEET_CELLS
        ):
            raise IntakeValidationError("A worksheet exceeds the dimension limits.")

    app_answers = _id_answer_map(workbook["1 App Basics"], answer_column=3)
    platform_answers = _id_answer_map(workbook["Platform + FinOps"], answer_column=5)
    application_name = _required_text(app_answers.get("A1"), "A1 application name")
    environments = _parse_environments(_required_text(app_answers.get("A3"), "A3 environments"))

    compute_units: list[ComputeUnit] = []
    gaps: list[Gap] = []
    for index, row in enumerate(
        workbook["2 Servers"].iter_rows(min_row=2, values_only=True),
        start=1,
    ):
        row = _pad_row(row, 11)
        if not any(value is not None for value in row):
            continue
        name = _required_text(row[0], "server name")
        runtime_text = _required_text(row[8], f"{name} runtime")
        runtime_hours = Decimal("730") if runtime_text.casefold() == "24x7" else None
        unit_id = f"compute-{index}"
        operating_system = _required_text(row[4], f"{name} operating system")
        license_model = _required_text(row[5], f"{name} license model")
        if runtime_hours is None:
            gaps.append(
                Gap(
                    id=f"{unit_id}.runtime",
                    unit_id=unit_id,
                    kind=GapKind.RUNTIME_HOURS,
                    prompt=f"How many hours per month does {name} run?",
                    reason="Server runtime is Unknown; cannot compute run-hours.",
                    raw_value=runtime_text,
                )
            )
        if "windows" in operating_system.casefold() and _license_requires_eligibility(
            license_model
        ):
            gaps.append(
                Gap(
                    id=f"{unit_id}.license",
                    unit_id=unit_id,
                    kind=GapKind.LICENSE_ELIGIBILITY,
                    license_product="Windows Server",
                    prompt=(
                        f"Confirm Software Assurance, Azure Hybrid Benefit, "
                        f"and license-vintage eligibility for {name}."
                    ),
                    reason=(
                        "Windows AHB and AWS BYOL are not equivalent and cannot be "
                        "priced safely from the license label alone."
                    ),
                    raw_value=license_model,
                )
            )
        count = _positive_int(row[2], f"{name} count")
        compute_units.append(
            ComputeUnit(
                id=unit_id,
                name=name,
                environment=_required_text(row[1], f"{name} environment"),
                count=count,
                operating_system=operating_system,
                license_model=license_model,
                vcpu_each=_positive_int(row[6], f"{name} vCPU"),
                ram_gb_each=_positive_decimal(row[7], f"{name} RAM"),
                runtime_hours_month=runtime_hours,
                topology_tier=TopologyTier(mode="single-zone", instance_count=count),
                storage_performance_profile=StoragePerformanceProfile(),
                network_transfer_profile=NetworkTransferProfile(
                    exclusion_policy_id=SHARED_EXCLUSION_POLICY
                ),
                backup_policy=BackupPolicy(
                    description="Covered by the approved shared-platform exclusion."
                ),
                support_tier=SupportTier(
                    name="Excluded",
                    exclusion_policy_id=SUPPORT_EXCLUSION_POLICY,
                ),
                observability_profile=ObservabilityProfile(
                    exclusion_policy_id=OBSERVABILITY_EXCLUSION_POLICY
                ),
                security_profile=SecurityProfile(
                    required_controls=[],
                    exclusion_policy_id=SECURITY_EXCLUSION_POLICY,
                ),
            )
        )

    database_units: list[DatabaseUnit] = []
    for index, row in enumerate(
        workbook["3 Databases"].iter_rows(min_row=2, values_only=True),
        start=1,
    ):
        row = _pad_row(row, 10)
        if not any(value is not None for value in row):
            continue
        name = _required_text(row[0], "database name")
        vcpu, ram_gb = _parse_vcpu_ram(_required_text(row[5], f"{name} vCPU and RAM"))
        topology_text = _required_text(row[7], f"{name} HA and backup")
        nodes = _positive_int(row[4], f"{name} nodes")
        topology_mode = _parse_database_topology(topology_text, nodes)
        engine = _required_text(row[2], f"{name} engine")
        license_model = _required_text(row[3], f"{name} license model")
        unit_id = f"database-{index}"
        if _is_sql_server(engine):
            gaps.append(
                Gap(
                    id=f"{unit_id}.license",
                    unit_id=unit_id,
                    kind=GapKind.LICENSE_ELIGIBILITY,
                    license_product="SQL Server",
                    prompt=(
                        f"Confirm SQL Server Software Assurance, License Mobility, "
                        f"AHB, and passive-secondary eligibility for {name}."
                    ),
                    reason=(
                        "SQL Server licensing differs between Azure AHB, AWS License "
                        "Mobility, and Dedicated Host treatment."
                    ),
                    raw_value={
                        "license_model": license_model,
                        "topology": topology_text,
                    },
                )
            )
        database_units.append(
            DatabaseUnit(
                id=unit_id,
                name=name,
                environment=_required_text(row[1], f"{name} environment"),
                engine=engine,
                license_model=license_model,
                nodes=nodes,
                vcpu_each=vcpu,
                ram_gb_each=ram_gb,
                size_gb=_positive_decimal(row[6], f"{name} size"),
                topology_tier=TopologyTier(
                    mode=topology_mode,
                    instance_count=nodes,
                ),
                storage_performance_profile=StoragePerformanceProfile(
                    capacity_gb=_positive_decimal(row[6], f"{name} size")
                ),
                network_transfer_profile=NetworkTransferProfile(
                    exclusion_policy_id=SHARED_EXCLUSION_POLICY
                ),
                backup_policy=BackupPolicy(description=topology_text),
                support_tier=SupportTier(
                    name="Excluded",
                    exclusion_policy_id=SUPPORT_EXCLUSION_POLICY,
                ),
                observability_profile=ObservabilityProfile(
                    exclusion_policy_id=OBSERVABILITY_EXCLUSION_POLICY
                ),
                security_profile=SecurityProfile(
                    required_controls=[],
                    exclusion_policy_id=SECURITY_EXCLUSION_POLICY,
                ),
            )
        )

    storage_units: list[StorageUnit] = []
    for index, row in enumerate(
        workbook["4 Storage"].iter_rows(min_row=2, values_only=True),
        start=1,
    ):
        row = _pad_row(row, 9)
        if not any(value is not None for value in row):
            continue
        name = _required_text(row[0], "storage name")
        performance_text = _sanitize_untrusted_text(row[7])
        target_iops, target_mbps = _parse_storage_performance(performance_text)
        unit_id = f"storage-{index}"
        if target_iops is None or target_mbps is None:
            gaps.append(
                Gap(
                    id=f"{unit_id}.performance",
                    unit_id=unit_id,
                    kind=GapKind.STORAGE_PERFORMANCE,
                    prompt=f"What IOPS and MB/s are required for {name}?",
                    reason="Storage performance is incomplete; IOPS and throughput cannot be priced.",
                    raw_value=performance_text or None,
                )
            )
        allocated_gb = _positive_decimal(row[4], f"{name} allocated GB")
        storage_units.append(
            StorageUnit(
                id=unit_id,
                name=name,
                environment=_required_text(row[1], f"{name} environment"),
                storage_type=_required_text(row[2], f"{name} type"),
                protocol=_required_text(row[3], f"{name} protocol"),
                allocated_gb=allocated_gb,
                used_gb=_nonnegative_decimal(row[5], f"{name} used GB"),
                storage_performance_profile=StoragePerformanceProfile(
                    capacity_gb=allocated_gb,
                    target_iops=target_iops,
                    target_mbps=target_mbps,
                ),
                backup_policy=BackupPolicy(
                    description=_required_text(row[6], f"{name} backup policy")
                ),
            )
        )

    raw_regions = _sanitize_untrusted_text(platform_answers.get("P1"))
    canonical_regions = _parse_supported_regions(raw_regions)
    region_gap = Gap(
        id="P1.regions",
        kind=GapKind.APPROVED_REGIONS,
        prompt="Which Azure and AWS regions are approved for this benchmark?",
        reason="Approved primary/DR regions not set; cannot select a PriceBook region.",
        raw_value=raw_regions or None,
    )
    if canonical_regions:
        region_gap.status = GapStatus.RESOLVED
        region_gap.canonical_value = canonical_regions
        region_gap.resolved_by = "Intake P1"
        region_gap.resolved_at = datetime.now(UTC)
    gaps.append(region_gap)

    if not compute_units and not database_units and not storage_units:
        raise IntakeValidationError(
            "The Intake must contain at least one compute, database, or storage workload."
        )

    return NormalizedIntake(
        application_name=application_name,
        environments=environments,
        compute_units=compute_units,
        database_units=database_units,
        storage_units=storage_units,
        gaps=gaps,
    )


def _id_answer_map(sheet: object, answer_column: int) -> dict[str, object]:
    answers: dict[str, object] = {}
    for row in sheet.iter_rows(min_row=2, values_only=True):
        row = _pad_row(row, answer_column)
        if row and row[0]:
            answers[_sanitize_untrusted_text(row[0])] = row[answer_column - 1]
    return answers


def _required_text(value: object, field: str) -> str:
    text = _sanitize_untrusted_text(value)
    if not text:
        raise IntakeValidationError(f"The Intake is missing {field}.")
    return text


def _parse_environments(value: str) -> list[str]:
    cleaned = value.replace(" and ", ", ")
    return [item.strip().rstrip(".") for item in cleaned.split(",") if item.strip()]


def _positive_int(value: object, field: str) -> int:
    try:
        parsed = Decimal(str(value))
    except Exception as error:
        raise IntakeValidationError(f"{field} must be a whole number.") from error
    if parsed != parsed.to_integral_value():
        raise IntakeValidationError(f"{field} must be a whole number.")
    if parsed <= 0:
        raise IntakeValidationError(f"{field} must be greater than zero.")
    return int(parsed)


def _positive_decimal(value: object, field: str) -> Decimal:
    parsed = _decimal(value, field)
    if parsed <= 0:
        raise IntakeValidationError(f"{field} must be greater than zero.")
    return parsed


def _nonnegative_decimal(value: object, field: str) -> Decimal:
    parsed = _decimal(value, field)
    if parsed < 0:
        raise IntakeValidationError(f"{field} must not be negative.")
    return parsed


def _decimal(value: object, field: str) -> Decimal:
    try:
        return Decimal(str(value))
    except Exception as error:
        raise IntakeValidationError(f"{field} must be numeric.") from error


def _parse_vcpu_ram(value: str) -> tuple[int, Decimal]:
    match = re.fullmatch(
        r"\s*(\d+)\s*vCPU\s*/\s*(\d+(?:\.\d+)?)\s*GB\s*",
        value,
        flags=re.IGNORECASE,
    )
    if not match:
        raise IntakeValidationError("Database vCPU and RAM must use '<vCPU> vCPU / <RAM> GB'.")
    return int(match.group(1)), Decimal(match.group(2))


def _parse_storage_performance(value: str) -> tuple[int | None, Decimal | None]:
    iops_match = re.search(r"(\d+)\s*IOPS", value, flags=re.IGNORECASE)
    throughput_match = re.search(r"(\d+(?:\.\d+)?)\s*MB/s", value, flags=re.IGNORECASE)
    return (
        int(iops_match.group(1)) if iops_match else None,
        Decimal(throughput_match.group(1)) if throughput_match else None,
    )


def _pad_row(row: tuple[object, ...], length: int) -> tuple[object, ...]:
    return row[:length] + (None,) * max(0, length - len(row))


def _sanitize_untrusted_text(value: object) -> str:
    text = str(value).strip() if value is not None else ""
    return f"'{text}" if text.startswith(("=", "+", "-", "@")) else text


def _license_requires_eligibility(value: str) -> bool:
    folded = re.sub(r"\s+", " ", value.strip().casefold())
    return any(
        term in folded
        for term in (
            "ahb",
            "azure hybrid benefit",
            "byol",
            "bring your own",
            "software assurance",
        )
    )


def _is_sql_server(value: str) -> bool:
    return re.search(r"(?<![a-z])sql\s+server(?![a-z])", value.casefold()) is not None


def _parse_supported_regions(value: str) -> dict[str, str] | None:
    folded = value.casefold()
    azure_match = re.search(r"(?<![a-z0-9-])eastus2(?![a-z0-9-])", folded)
    aws_match = re.search(r"(?<![a-z0-9-])us-east-1(?![a-z0-9-])", folded)
    if azure_match and aws_match:
        return {"azure_region": "eastus2", "aws_region": "us-east-1"}
    return None


def _parse_database_topology(value: str, nodes: int) -> str:
    folded = value.casefold()
    explicit_no_ha = bool(re.search(r"\bno\s+ha\b|\bwithout\s+ha\b", folded))
    explicit_ha = bool(
        re.search(r"\bha\b|\breplica\b|\bmulti-az\b|\bzone[- ]redundant\b", folded)
    ) and not explicit_no_ha
    if explicit_no_ha:
        if nodes != 1:
            raise IntakeValidationError("A no-HA database must declare exactly one node.")
        return "single-zone"
    if explicit_ha:
        if nodes < 2:
            raise IntakeValidationError("An HA database must declare at least two nodes.")
        return "zone-redundant"
    if nodes > 1:
        raise IntakeValidationError("Multiple database nodes require an explicit HA topology.")
    return "single-zone"


def _local_name(tag: object) -> str:
    return str(tag).rsplit("}", 1)[-1].casefold()


def _parse_part(xml_bytes: bytes):
    # defusedxml honours the declared encoding (UTF-16 included) and refuses DTDs and entities.
    try:
        return DefusedElementTree.fromstring(xml_bytes, forbid_dtd=True)
    except (ParseError, ValueError) as error:
        raise IntakeValidationError("The OpenXML package contains invalid XML.") from error


def _validate_workbook_location(content_types: bytes, package_rels: bytes) -> None:
    """openpyxl finds the workbook through [Content_Types].xml, so pin it and every sheet part in place."""
    for override in _parse_part(content_types).iter():
        if _local_name(override.tag) != "override":
            continue
        part = override.attrib.get("PartName", "").casefold()
        content_type = override.attrib.get("ContentType", "").casefold()
        if content_type.endswith(".main+xml") and part != "/xl/workbook.xml":
            raise IntakeValidationError("The workbook part is not at xl/workbook.xml and is not accepted.")
        if content_type == WORKSHEET_CONTENT_TYPE and not (
            part.startswith("/xl/worksheets/") and part.endswith(".xml") and "/" not in part[15:]
        ):
            raise IntakeValidationError("A worksheet is stored outside xl/worksheets/ and is not accepted.")
        if content_type.endswith(("chartsheet+xml", "dialogsheet+xml", "macrosheet+xml")):
            raise IntakeValidationError("Chart, dialog, and macro sheets are not accepted.")
    for relationship in _parse_part(package_rels).iter():
        if _local_name(relationship.tag) != "relationship":
            continue
        if relationship.attrib.get("Type", "").casefold().endswith(OFFICE_DOCUMENT_RELATIONSHIP):
            target = posixpath.normpath(relationship.attrib.get("Target", "").lstrip("/"))
            if target.casefold() != "xl/workbook.xml":
                raise IntakeValidationError("The workbook part is not at xl/workbook.xml and is not accepted.")


def _is_range_reference(text: str) -> bool:
    """True only for comma-separated Sheet!range references, with Excel's '' quote escaping."""
    if not text or len(text) > MAX_DEFINED_NAME_LENGTH:
        return False
    index = 0
    while True:
        if text.startswith("'", index):
            index += 1
            length = 0
            while True:
                if index >= len(text):
                    return False
                if text[index] == "'":
                    if text.startswith("''", index):
                        index += 2
                        length += 1
                        continue
                    index += 1
                    break
                index += 1
                length += 1
            if not 1 <= length <= 31:
                return False
        else:
            sheet = UNQUOTED_SHEET.match(text, index)
            if not sheet:
                return False
            index = sheet.end()
        if not text.startswith("!", index):
            return False
        cells = CELL_RANGE.match(text, index + 1)
        if not cells:
            return False
        index = cells.end()
        if index == len(text):
            return True
        if text[index] != ",":
            return False
        index += 1


def _validate_workbook_sheets(workbook_bytes: bytes, rels_bytes: bytes) -> None:
    """Every <sheet> must load a worksheet relationship, and defined names can't carry formulas."""
    types = {
        relationship.attrib.get("Id"): relationship.attrib.get("Type", "").casefold()
        for relationship in _parse_part(rels_bytes).iter()
        if _local_name(relationship.tag) == "relationship"
    }
    for element in _parse_part(workbook_bytes).iter():
        name = _local_name(element.tag)
        if name == "sheet" and not types.get(element.attrib.get(RELATIONSHIP_ID), "").endswith(WORKSHEET_RELATIONSHIP):
            raise IntakeValidationError("A sheet in the workbook is not a worksheet and is not accepted.")
        if name == "definedname" and not (
            element.attrib.get("name", "").casefold() in SAFE_DEFINED_NAMES
            and _is_range_reference((element.text or "").strip())
        ):
            raise IntakeValidationError("Workbook formulas are not accepted in Intake fields (defined names).")


def _validate_worksheet_targets(rels_bytes: bytes) -> None:
    """Every worksheet the workbook loads must be a checked xl/worksheets/*.xml part."""
    for relationship in _parse_part(rels_bytes).iter():
        if _local_name(relationship.tag) != "relationship":
            continue
        if not relationship.attrib.get("Type", "").casefold().endswith(WORKSHEET_RELATIONSHIP):
            continue
        target = relationship.attrib.get("Target", "").replace("\\", "/")
        resolved = posixpath.normpath(target.lstrip("/") if target.startswith("/") else f"xl/{target}")
        folded = resolved.casefold()
        if not (folded.startswith("xl/worksheets/") and folded.endswith(".xml") and "/" not in folded[14:]):
            raise IntakeValidationError("A worksheet is stored outside xl/worksheets/ and is not accepted.")


def _validate_worksheet_xml(xml_bytes: bytes) -> None:
    try:
        root = DefusedElementTree.fromstring(xml_bytes)
    except (ParseError, ValueError) as error:
        raise IntakeValidationError("A worksheet contains invalid XML.") from error

    dimension = next(
        (element for element in root.iter() if element.tag.rsplit("}", 1)[-1] == "dimension"),
        None,
    )
    reference = dimension.attrib.get("ref") if dimension is not None else None
    if not reference:
        raise IntakeValidationError("A worksheet does not declare a bounded dimension.")
    try:
        min_column, min_row, max_column, max_row = range_boundaries(reference)
    except ValueError as error:
        raise IntakeValidationError("A worksheet declares an invalid dimension.") from error
    if (
        min_row < 1
        or min_column < 1
        or max_row > MAX_SHEET_ROWS
        or max_column > MAX_SHEET_COLUMNS
        or max_row * max_column > MAX_SHEET_CELLS
    ):
        raise IntakeValidationError("A worksheet exceeds the dimension limits.")

    actual_cells = 0
    actual_max_row = 0
    actual_max_column = 0
    for element in root.iter():
        if element.tag.rsplit("}", 1)[-1] != "c":
            continue
        coordinate = element.attrib.get("r")
        if not coordinate:
            raise IntakeValidationError("A worksheet cell is missing its coordinate.")
        try:
            row, column = coordinate_to_tuple(coordinate)
        except ValueError as error:
            raise IntakeValidationError("A worksheet cell has an invalid coordinate.") from error
        actual_cells += 1
        actual_max_row = max(actual_max_row, row)
        actual_max_column = max(actual_max_column, column)
        if (
            actual_cells > MAX_SHEET_CELLS
            or actual_max_row > MAX_SHEET_ROWS
            or actual_max_column > MAX_SHEET_COLUMNS
        ):
            raise IntakeValidationError("A worksheet exceeds the dimension limits.")
    if actual_max_row > max_row or actual_max_column > max_column:
        raise IntakeValidationError(
            "A worksheet contains cells outside its declared dimension."
        )


def _validate_relationships(member_name: str, xml_bytes: bytes) -> None:
    try:
        root = DefusedElementTree.fromstring(xml_bytes)
    except (ParseError, ValueError) as error:
        raise IntakeValidationError("The OpenXML package contains invalid relationships.") from error

    folded_name = member_name.casefold()
    if folded_name == "_rels/.rels":
        base_path = ""
    else:
        marker = "/_rels/"
        marker_index = folded_name.rfind(marker)
        if marker_index < 0:
            raise IntakeValidationError("The OpenXML package contains invalid relationships.")
        source_name = (
            member_name[:marker_index]
            + "/"
            + member_name[marker_index + len(marker) : -len(".rels")]
        )
        base_path = posixpath.dirname(source_name)

    for relationship in root.iter():
        if relationship.tag.rsplit("}", 1)[-1] != "Relationship":
            continue
        if relationship.attrib.get("TargetMode", "").casefold() == "external":
            raise IntakeValidationError("External workbook relationships are not accepted.")
        target = relationship.attrib.get("Target", "")
        if not target or "://" in target or target.startswith(("\\", "//")):
            raise IntakeValidationError("The OpenXML package contains an unsafe relationship.")
        normalized_target = target.lstrip("/") if target.startswith("/") else target
        relationship_base = "" if target.startswith("/") else base_path
        resolved = posixpath.normpath(
            posixpath.join(relationship_base, normalized_target.replace("\\", "/"))
        )
        if resolved == ".." or resolved.startswith("../"):
            raise IntakeValidationError("The OpenXML package contains an unsafe relationship.")


sanitize_untrusted_text = _sanitize_untrusted_text
