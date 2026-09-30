import io
import re
from decimal import Decimal
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import load_workbook

from app.models import ApplicationDetail, Cloud
from app.workbook_export import build_priced_workbook


def _upload(client: object, fixture: Path) -> dict:
    with fixture.open("rb") as workbook:
        response = client.post(
            "/api/intakes",
            files={
                "file": (
                    fixture.name,
                    workbook,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
    assert response.status_code == 201, response.text
    return response.json()


def _resolve_demo_gaps(client: object, application_id: str) -> dict:
    client.post(
        f"/api/applications/{application_id}/gaps/P1.regions/resolve",
        json={
            "resolved_by": "FinOps Reviewer",
            "azure_region": "eastus2",
            "aws_region": "us-east-1",
        },
    ).raise_for_status()
    response = client.post(
        f"/api/applications/{application_id}/gaps/compute-3.runtime/resolve",
        json={
            "resolved_by": "Application Owner",
            "runtime_hours_month": "220",
        },
    )
    response.raise_for_status()
    return response.json()


def test_gap_resolution_unlocks_review_baseline_and_export(
    client: object,
    repository_root: Path,
) -> None:
    detail = _upload(client, repository_root / "samples" / "synthetic_intake_completed.xlsx")
    application_id = detail["application"]["id"]
    assert detail["comparison"]["state"] == "DraftBenchmark"
    assert detail["comparison"]["aws_monthly_total"] is None
    assert detail["comparison"]["can_export"] is False
    assert client.get(
        f"/api/applications/{application_id}/comparison/export"
    ).status_code == 409

    region_response = client.post(
        f"/api/applications/{application_id}/gaps/P1.regions/resolve",
        json={
            "resolved_by": "FinOps Reviewer",
            "azure_region": "eastus2",
            "aws_region": "us-east-1",
        },
    )
    assert region_response.status_code == 200
    provisional = region_response.json()["comparison"]
    assert provisional["state"] == "DraftBenchmark"
    assert provisional["aws_monthly_total"] is not None
    assert any(line["status"] == "Unpriced" for line in provisional["line_items"])
    provisional_scenarios = {
        item["id"]: item for item in provisional["commercial_scenarios"]
    }
    assert provisional_scenarios["aws-list"]["available"] is True
    assert provisional_scenarios["aws-savings-plan"]["available"] is False
    assert provisional_scenarios["azure-reservation"]["available"] is False
    assert all(
        not item["available"]
        and item["reference_provider"] is None
        and item["target_provider"] is None
        and item["aws_public_monthly_total"] is None
        and item["azure_public_monthly_total"] is None
        and item["points"] == []
        and "ReviewBaseline" in item["unavailable_reason"]
        for item in provisional["breakeven_sensitivities"]
    )

    runtime_response = client.post(
        f"/api/applications/{application_id}/gaps/compute-3.runtime/resolve",
        json={
            "resolved_by": "Application Owner",
            "runtime_hours_month": "220",
        },
    )
    assert runtime_response.status_code == 200
    baseline = runtime_response.json()["comparison"]
    assert baseline["state"] == "ReviewBaseline"
    assert baseline["can_export"] is True
    scenarios = {item["id"]: item for item in baseline["commercial_scenarios"]}
    assert scenarios["aws-list"]["monthly_total"] == baseline["aws_monthly_total"]
    assert scenarios["azure-list"]["monthly_total"] == baseline["azure_monthly_total"]
    assert scenarios["aws-savings-plan"]["monthly_total"] == "2660.43"
    assert scenarios["aws-savings-plan"]["covered_monthly_total"] == "474.06"
    assert scenarios["aws-savings-plan"]["uncovered_monthly_total"] == "2186.37"
    assert scenarios["aws-reservation"]["monthly_total"] == "2593.09"
    assert scenarios["azure-savings-plan"]["monthly_total"] == "2681.55"
    assert scenarios["azure-reservation"]["monthly_total"] == "2609.14"
    assert scenarios["azure-reservation"]["covered_monthly_total"] == "401.73"
    assert scenarios["azure-reservation"]["uncovered_monthly_total"] == "2207.41"
    assert all(item["available"] for item in scenarios.values())
    assert baseline["presentation"]["show_aws"] is False
    sensitivities = {
        item["id"]: item for item in baseline["breakeven_sensitivities"]
    }
    assert set(sensitivities) == {
        "list-parity",
        "savingsplan-parity",
        "reservation-parity",
    }
    assert all(item["available"] for item in sensitivities.values())
    assert sensitivities["list-parity"]["reference_provider"] == "AWS"
    assert sensitivities["list-parity"]["target_provider"] == "Azure"
    assert sensitivities["list-parity"]["points"] == [
        {
            "reference_discount_percent": "0",
            "target_discount_to_parity_percent": "0.68",
            "additional_discount_advantage_percent": "0.68",
        },
        {
            "reference_discount_percent": "10",
            "target_discount_to_parity_percent": "10.61",
            "additional_discount_advantage_percent": "0.61",
        },
        {
            "reference_discount_percent": "20",
            "target_discount_to_parity_percent": "20.55",
            "additional_discount_advantage_percent": "0.55",
        },
    ]
    assert sensitivities["savingsplan-parity"]["points"][0][
        "target_discount_to_parity_percent"
    ] == "0.79"
    assert sensitivities["reservation-parity"]["points"][0][
        "target_discount_to_parity_percent"
    ] == "0.62"
    assert all(
        item["public_price_only"]
        and not item["actual_contracted_price"]
        and item["workload_isolated"]
        and "not an EA" in item["disclosure"]
        for item in sensitivities.values()
    )
    assert baseline["assumptions"]["commitment_term_years"] == 3
    assert baseline["assumptions"]["commitment_payment_option"] == "All Upfront"
    assert baseline["assumptions"]["commitment_utilization_percent"] == "100"
    assert baseline["assumptions"]["commitment_tenancy"] == "Shared"
    assert any("26,280 hours" in warning for warning in baseline["warnings"])
    assert any("/savingsPlan/v1.0/" in url for url in baseline["source_urls"])
    assert baseline["headline_delta"] is not None
    assert any(
        line["azure_amount"] is not None
        and line["aws_amount"] is not None
        and float(line["azure_amount"]) > float(line["aws_amount"])
        for line in baseline["line_items"]
    )
    assert any(
        line["component"] == "RHEL subscription uplift"
        for line in baseline["line_items"]
    )
    assert any(
        line["component"] == "PostgreSQL HA standby"
        for line in baseline["line_items"]
    )
    priced_line = next(
        line for line in baseline["line_items"] if line["status"] == "Priced"
    )
    assert priced_line["aws_commercial"]["list_rate"] == priced_line["aws_rate"]
    assert priced_line["aws_commercial"]["list_amount"] == priced_line["aws_amount"]
    assert (
        priced_line["azure_commercial"]["list_rate"] == priced_line["azure_rate"]
    )
    assert (
        priced_line["azure_commercial"]["list_amount"] == priced_line["azure_amount"]
    )
    windows_line = next(
        line
        for line in baseline["line_items"]
        if line["component"] == "Windows license uplift"
    )
    assert windows_line["aws_commercial"]["savings_plan_rate"] == windows_line["aws_rate"]
    assert (
        windows_line["azure_commercial"]["reservation_rate"]
        == windows_line["azure_rate"]
    )
    storage_line = next(
        line
        for line in baseline["line_items"]
        if line["component"] == "Block storage capacity"
    )
    assert (
        storage_line["azure_commercial"]["savings_plan_amount"]
        == storage_line["azure_commercial"]["list_amount"]
    )
    part_time_compute = next(
        line
        for line in baseline["line_items"]
        if line["unit_name"] == "app-test" and line["component"] == "VM compute"
    )
    assert part_time_compute["aws_quantity"] == "220"
    assert part_time_compute["aws_commercial"]["reservation_quantity"] == "730"
    assert part_time_compute["aws_commercial"]["reservation_amount"] == "58.11"
    assert part_time_compute["azure_commercial"]["savings_plan_amount"] == "67.83"
    web_prod_lines = [
        line
        for line in baseline["line_items"]
        if line["unit_name"] == "web-prod"
    ]
    assert {line["component"] for line in web_prod_lines} == {
        "VM compute",
        "Windows license uplift",
    }
    assert sum(Decimal(line["aws_amount"]) for line in web_prod_lines) == Decimal("562.98")
    assert sum(Decimal(line["azure_amount"]) for line in web_prod_lines) == Decimal("563.56")

    export_response = client.get(
        f"/api/applications/{application_id}/comparison/export"
    )
    assert export_response.status_code == 200
    exported = export_response.json()
    assert exported["comparison"]["run_hash"] == baseline["run_hash"]
    assert exported["comparison"]["product_label"] == "public-list run-rate benchmark"
    assert exported["comparison"]["skumap_approval"]["approver"] == "Sample Approver (demo data)"
    assert exported["comparison"]["aws_monthly_total"] == baseline["aws_monthly_total"]

    workbook_response = client.get(
        f"/api/applications/{application_id}/comparison/workbook"
    )
    assert workbook_response.status_code == 200
    assert workbook_response.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert workbook_response.headers["content-disposition"].startswith(
        f'attachment; filename="priced-{application_id}.xlsx"; filename*=UTF-8'
    )
    workbook = load_workbook(
        filename=io.BytesIO(workbook_response.content),
        data_only=False,
    )
    assert workbook["1 App Basics"]["C2"].value == "StatementHub (synthetic)"
    assert workbook.sheetnames[:5] == [
        "Pricing Summary",
        "Pricing Detail",
        "Pricing Assumptions",
        "Pricing Exclusions",
        "Pricing Evidence",
    ]
    assert workbook["Pricing Summary"]["B9"].value.startswith("=ROUND(SUM(")
    assert workbook["Pricing Summary"]["B10"].value.startswith("=ROUND(SUM(")
    assert Decimal(str(workbook["Pricing Summary"]["C9"].value)) == Decimal(
        scenarios["azure-list"]["monthly_total"]
    )
    assert workbook["Pricing Summary"]["D9"].value.startswith("=IF(")
    assert Decimal(str(workbook["Pricing Summary"]["C10"].value)) == Decimal(
        scenarios["azure-savings-plan"]["monthly_total"]
    )
    assert Decimal(str(workbook["Pricing Summary"]["C11"].value)) == Decimal(
        scenarios["azure-reservation"]["monthly_total"]
    )
    assert workbook["Pricing Summary"]["E10"].value.startswith("=ROUND(SUM(")
    assert workbook["Pricing Summary"]["F10"].value.startswith("=ROUND(SUM(")
    assert workbook["Pricing Evidence"]["B4"].value == baseline["pricebook_snapshot_id"]
    assert any(
        "prices.azure.com/api/retail/prices" in str(cell.value)
        for row in workbook["Pricing Evidence"].iter_rows()
        for cell in row
    )
    evidence_row_ids = {
        workbook["Pricing Evidence"].cell(row=row, column=5).value
        for row in range(1, workbook["Pricing Evidence"].max_row + 1)
    }
    assert any(
        value
        and "6eeeab906a8dafe2858f4e29653c61ce8849dd8bc35361f65c75ddfc21c3d5f1"
        in value
        for value in evidence_row_ids
    )
    assert workbook["Pricing Evidence"]["B6"].value == "trust-20260922e"
    evidence_values = {
        workbook["Pricing Evidence"].cell(row=row, column=1).value:
        workbook["Pricing Evidence"].cell(row=row, column=2).value
        for row in range(1, workbook["Pricing Evidence"].max_row + 1)
    }
    assert evidence_values["Run hash"] == baseline["run_hash"]
    assumption_values = {
        workbook["Pricing Assumptions"].cell(row=row, column=1).value:
        workbook["Pricing Assumptions"].cell(row=row, column=2).value
        for row in range(1, workbook["Pricing Assumptions"].max_row + 1)
    }
    assert assumption_values["Commitment term"] == "3 years"
    assert assumption_values["Commitment payment option"] == "All Upfront"
    assert assumption_values["Commitment utilization"] == "100%"
    assert assumption_values["Commitment tenancy"] == "Shared"
    assert workbook["Pricing Detail"]["H6"].value.startswith("=IF(")
    assert workbook["Pricing Detail"]["K6"].value.startswith("=IF(")
    assert workbook["Pricing Detail"]["N6"].value.startswith("=IF(")
    generated_sheet_text = {
        str(cell.value)
        for sheet_name in workbook.sheetnames[:5]
        for row in workbook[sheet_name].iter_rows()
        for cell in row
        if cell.value is not None
    }
    assert any(
        "26,280 hours" in value
        and "730 hours per instance-month" in value
        for value in generated_sheet_text
    )
    assert not any(
        "aws" in value.casefold() or "us-east" in value.casefold()
        for value in generated_sheet_text
    )
    exclusion_categories = {
        workbook["Pricing Exclusions"].cell(row=row, column=1).value
        for row in range(5, workbook["Pricing Exclusions"].max_row + 1)
    }
    assert "Support" in exclusion_categories
    assert "Observability" in exclusion_categories
    assert any(
        source["source_type"] == "DemoAssumption"
        for source in baseline["rate_sources"]
    )


def test_workbook_export_preserves_colliding_source_sheet(
    client: object,
    repository_root: Path,
    tmp_path: Path,
) -> None:
    source = repository_root / "samples" / "synthetic_intake_completed.xlsx"
    workbook = load_workbook(source)
    workbook.create_sheet("Pricing Summary")["A1"] = "SOURCE MARKER"
    path = tmp_path / "sheet-collision.xlsx"
    workbook.save(path)

    detail = _upload(client, path)
    application_id = detail["application"]["id"]
    _resolve_demo_gaps(client, application_id)
    response = client.get(
        f"/api/applications/{application_id}/comparison/workbook"
    )
    assert response.status_code == 200
    exported = load_workbook(io.BytesIO(response.content), data_only=False)
    assert exported["Pricing Summary"]["A1"].value == "SOURCE MARKER"
    assert exported["Pricing Summary (2)"]["A1"].value == "Azure Pricing Summary"


def test_aws_presentation_flag_restores_full_workbook(
    client: object,
    repository_root: Path,
    monkeypatch: object,
) -> None:
    monkeypatch.setenv("PRESENT_AWS_PRICING", "true")
    detail = _upload(
        client,
        repository_root / "samples" / "synthetic_intake_completed.xlsx",
    )
    application_id = detail["application"]["id"]
    resolved = _resolve_demo_gaps(client, application_id)
    comparison = resolved["comparison"]
    assert comparison["presentation"]["show_aws"] is True
    scenarios = {scenario["id"]: scenario for scenario in comparison["commercial_scenarios"]}
    assert set(scenarios) == {
        "azure-list",
        "azure-savings-plan",
        "azure-reservation",
        "aws-list",
        "aws-savings-plan",
        "aws-reservation",
    }

    response = client.get(
        f"/api/applications/{application_id}/comparison/workbook"
    )
    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.content), data_only=False)
    summary = workbook["Pricing Summary"]
    pricing_detail = workbook["Pricing Detail"]
    assert workbook.sheetnames[:5] == [
        "Pricing Summary",
        "Pricing Detail",
        "Pricing Assumptions",
        "Pricing Exclusions",
        "Pricing Evidence",
    ]
    assert workbook["1 App Basics"]["C2"].value == "StatementHub (synthetic)"
    assert summary["A1"].value == "Cloud Pricing Summary"
    assert pricing_detail["A1"].value == "Cloud Pricing Detail"
    assert [pricing_detail.cell(row=5, column=column).value for column in range(16, 26)] == [
        "AWS quantity",
        "AWS quantity unit",
        "AWS List rate",
        "AWS List monthly",
        "AWS Savings Plan covered",
        "AWS Savings Plan uncovered",
        "AWS Savings Plan monthly",
        "AWS RI covered",
        "AWS RI uncovered",
        "AWS RI monthly",
    ]
    assert pricing_detail["G5"].value == "Azure List rate"
    assert pricing_detail["K5"].value == "Azure Savings Plan monthly"
    assert pricing_detail["N5"].value == "Azure Reservation monthly"

    detail_last_row = 5 + len(comparison["line_items"])
    scenario_rows = {
        "azure-list": (9, "H", None, None),
        "azure-savings-plan": (10, "K", "I", "J"),
        "azure-reservation": (11, "N", "L", "M"),
        "aws-list": (12, "S", None, None),
        "aws-savings-plan": (13, "V", "T", "U"),
        "aws-reservation": (14, "Y", "W", "X"),
    }
    for scenario_id, (row, total, covered, uncovered) in scenario_rows.items():
        scenario = scenarios[scenario_id]
        assert scenario["available"] is True
        assert summary[f"A{row}"].value == scenario["label"]
        assert Decimal(str(summary[f"C{row}"].value)) == Decimal(scenario["monthly_total"])
        assert summary[f"B{row}"].value == (
            f"=ROUND(SUM('Pricing Detail'!{total}6:{total}{detail_last_row}),2)"
        )
        basis_check = {
            "I": "AH", "L": "AI", "T": "AJ", "W": "AK",
        }.get(covered)
        assert summary[f"D{row}"].value == (
            f'=IF(OR(ABS(B{row}-C{row})>=0.005,'
            f'COUNTIF(\'Pricing Detail\'!{basis_check}6:{basis_check}{detail_last_row},'
            f'"MISMATCH")>0),"MISMATCH","Reconciled")'
            if basis_check else
            f'=IF(ABS(B{row}-C{row})<0.005,"Reconciled","MISMATCH")'
        )
        if covered:
            assert summary[f"E{row}"].value == (
                f"=ROUND(SUM('Pricing Detail'!{covered}6:{covered}{detail_last_row}),2)"
            )
            assert summary[f"F{row}"].value == (
                f"=ROUND(SUM('Pricing Detail'!{uncovered}6:{uncovered}{detail_last_row}),2)"
            )
            assert scenario["covered_monthly_total"] is not None
            assert scenario["uncovered_monthly_total"] is not None
        else:
            assert summary[f"E{row}"].value == 0
            assert summary[f"F{row}"].value == f"=B{row}"
        provider = scenario_id.split("-", 1)[0]
        offer = scenario_id.split("-", 1)[1].replace("-", "_")
        amounts = [
            Decimal(line[f"{provider}_commercial"][f"{offer}_amount"])
            for line in comparison["line_items"]
            if line["status"] == "Priced"
        ]
        assert sum(amounts, Decimal("0")) == Decimal(scenario["monthly_total"])
        if covered:
            for part in ("covered", "uncovered"):
                line_total = sum(
                    (
                        Decimal(
                            line[f"{provider}_commercial"][f"{offer}_{part}_amount"]
                        )
                        for line in comparison["line_items"]
                        if line["status"] == "Priced"
                    ),
                    Decimal("0"),
                )
                assert line_total == Decimal(scenario[f"{part}_monthly_total"])

    for row, line in enumerate(comparison["line_items"], start=6):
        for column, provider, offer in (
            ("I", "azure", "savings_plan_covered"),
            ("J", "azure", "savings_plan_uncovered"),
            ("L", "azure", "reservation_covered"),
            ("M", "azure", "reservation_uncovered"),
            ("T", "aws", "savings_plan_covered"),
            ("U", "aws", "savings_plan_uncovered"),
            ("W", "aws", "reservation_covered"),
            ("X", "aws", "reservation_uncovered"),
        ):
            amount = line[f"{provider}_commercial"][f"{offer}_amount"]
            value = pricing_detail[f"{column}{row}"].value
            assert (None if value is None else Decimal(str(value))) == (
                None if amount is None else Decimal(amount)
            )
        for total, quantity, rate in (("H", "E", "G"), ("S", "P", "R")):
            assert pricing_detail[f"{total}{row}"].value == (
                f'=IF(OR(C{row}<>"Priced",{quantity}{row}="",'
                f'{rate}{row}=""),"",ROUND({quantity}{row}*{rate}{row},2))'
            )
        if line["demo_assumption"]:
            assert pricing_detail[f"G{row}"].fill.fgColor.rgb.endswith("FFF2CC")
            for column in ("K", "N"):
                assert pricing_detail[f"{column}{row}"].fill.fgColor.rgb.endswith("FFF2CC")
            for column in ("R", "S", "T", "U", "W", "X"):
                assert pricing_detail[f"{column}{row}"].fill.patternType is None
        for total, covered, uncovered in (
            ("K", "I", "J"), ("N", "L", "M"),
            ("V", "T", "U"), ("Y", "W", "X"),
        ):
            assert pricing_detail[f"{total}{row}"].value == (
                f'=IF(COUNT({covered}{row}:{uncovered}{row})=2,'
                f'ROUND(SUM({covered}{row}:{uncovered}{row}),2),"")'
            )
        for provider, offer, rate_column, quantity_column, check_column in (
            ("azure", "savings_plan", "Z", "AA", "AH"),
            ("azure", "reservation", "AB", "AC", "AI"),
            ("aws", "savings_plan", "AD", "AE", "AJ"),
            ("aws", "reservation", "AF", "AG", "AK"),
        ):
            commercial = line[f"{provider}_commercial"]
            covered_amount = commercial[f"{offer}_covered_amount"]
            if covered_amount is not None and Decimal(covered_amount) > 0:
                assert Decimal(str(pricing_detail[f"{rate_column}{row}"].value)) == Decimal(
                    commercial[f"{offer}_rate"]
                )
                assert Decimal(str(pricing_detail[f"{quantity_column}{row}"].value)) == Decimal(
                    commercial[f"{offer}_quantity"]
                )
                assert pricing_detail[f"{check_column}{row}"].value.startswith(
                    f'=IF(ABS(ROUND({rate_column}{row}*{quantity_column}{row},2)'
                )
            else:
                assert pricing_detail[f"{rate_column}{row}"].value is None
                assert pricing_detail[f"{quantity_column}{row}"].value is None
                assert pricing_detail[f"{check_column}{row}"].value is None

    assert summary["A17"].value == "Public basis"
    assert summary["C17"].value == "Target discount to parity"
    assert "Azure RHEL demo assumptions affect modeled parity" in summary["E18"].value
    for start_row, sensitivity in zip(
        range(18, 27, 3),
        comparison["breakeven_sensitivities"],
        strict=True,
    ):
        for row, point in enumerate(sensitivity["points"], start=start_row):
            assert summary[f"A{row}"].value == sensitivity["label"]
            assert summary[f"F{row}"].value == (
                sensitivity["target_provider"] or "Equal"
            )
            assert Decimal(str(summary[f"C{row}"].value)) * 100 == Decimal(
                point["target_discount_to_parity_percent"]
            )
    assert not any(
        "Lower-cost provider" == cell.value
        for row in summary.iter_rows()
        for cell in row
    )
    summary_notes = {
        cell.value for row in summary.iter_rows() for cell in row
        if isinstance(cell.value, str)
    }
    assert set(comparison["warnings"]) <= summary_notes
    assert workbook["Pricing Assumptions"]["A12"].value == "AWS region"
    assert any(
        "aws" == str(cell.value).casefold()
        for row in workbook["Pricing Evidence"].iter_rows()
        for cell in row
    )


def test_breakeven_aws_target_is_explicit_only_in_full_presentation(
    client: object,
    repository_root: Path,
) -> None:
    source = (repository_root / "samples" / "synthetic_intake_completed.xlsx").read_bytes()
    uploaded = _upload(
        client, repository_root / "samples" / "synthetic_intake_completed.xlsx"
    )
    resolved = _resolve_demo_gaps(client, uploaded["application"]["id"])
    detail = ApplicationDetail.model_validate(resolved)
    sensitivity = detail.comparison.breakeven_sensitivities[0]
    sensitivity.reference_provider = Cloud.AZURE
    sensitivity.target_provider = Cloud.AWS
    sensitivity.points[0].target_discount_to_parity_percent = Decimal("12.34")
    sensitivity.points[0].additional_discount_advantage_percent = Decimal("2.34")
    warning = "AWS public pricing for approved East US 2 and us-east-1 mappings."
    detail.comparison.warnings.append(warning)

    detail.comparison.presentation.show_aws = True
    full = load_workbook(io.BytesIO(build_priced_workbook(source, detail)))
    full_summary = full["Pricing Summary"]
    assert full_summary["C17"].value == "Target discount to parity"
    assert full_summary["F17"].value == "Target provider"
    assert full_summary["F18"].value == "AWS"
    assert Decimal(str(full_summary["C18"].value)) == Decimal("0.1234")
    assert Decimal(str(full_summary["D18"].value)) == Decimal("0.0234")
    assert warning in {
        cell.value for row in full_summary.iter_rows() for cell in row
        if isinstance(cell.value, str)
    }

    detail.comparison.presentation.show_aws = False
    azure_only = load_workbook(io.BytesIO(build_priced_workbook(source, detail)))
    azure_summary = azure_only["Pricing Summary"]
    assert azure_summary["C14"].value == "Azure discount to parity"
    assert azure_summary["C15"].value == "Already at or below parity"
    assert azure_summary["D15"].value == 0
    assert azure_summary["F14"].value is None
    assert warning not in {
        cell.value for row in azure_summary.iter_rows() for cell in row
        if isinstance(cell.value, str)
    }


def test_workbook_export_handles_case_insensitive_sheet_collisions(
    client: object,
    repository_root: Path,
    tmp_path: Path,
) -> None:
    source = repository_root / "samples" / "synthetic_intake_completed.xlsx"
    workbook = load_workbook(source)
    workbook.create_sheet("pricing detail")["A1"] = "SOURCE DETAIL"
    path = tmp_path / "case-collision.xlsx"
    workbook.save(path)

    detail = _upload(client, path)
    application_id = detail["application"]["id"]
    _resolve_demo_gaps(client, application_id)
    response = client.get(
        f"/api/applications/{application_id}/comparison/workbook"
    )
    assert response.status_code == 200
    exported = load_workbook(io.BytesIO(response.content), data_only=False)
    assert exported["pricing detail"]["A1"].value == "SOURCE DETAIL"
    generated = next(
        sheet
        for sheet in exported.worksheets
        if sheet.title.casefold().startswith("pricing detail (")
    )
    assert generated["A1"].value == "Azure Pricing Detail"
    assert generated.title in exported["Pricing Summary"]["B9"].value


def test_workbook_without_calc_properties_exports_successfully(
    client: object,
    repository_root: Path,
) -> None:
    source = repository_root / "samples" / "synthetic_intake_completed.xlsx"
    stream = io.BytesIO()
    with ZipFile(source) as original, ZipFile(stream, "w", ZIP_DEFLATED) as modified:
        for member in original.infolist():
            content = original.read(member.filename)
            if member.filename == "xl/workbook.xml":
                content = re.sub(rb"<calcPr\b[^>]*/>", b"", content)
            modified.writestr(member, content)

    response = client.post(
        "/api/intakes",
        files={
            "file": (
                "without-calc-properties.xlsx",
                stream.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert response.status_code == 201
    application_id = response.json()["application"]["id"]
    _resolve_demo_gaps(client, application_id)
    exported = client.get(
        f"/api/applications/{application_id}/comparison/workbook"
    )
    assert exported.status_code == 200
    workbook = load_workbook(io.BytesIO(exported.content), data_only=False)
    assert workbook.calculation is not None


def test_unicode_filename_uses_safe_content_disposition(
    client: object,
    repository_root: Path,
) -> None:
    source = repository_root / "samples" / "synthetic_intake_completed.xlsx"
    response = client.post(
        "/api/intakes",
        files={
            "file": (
                "アプリ台帳.xlsx",
                source.read_bytes(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert response.status_code == 201
    application_id = response.json()["application"]["id"]
    _resolve_demo_gaps(client, application_id)
    exported = client.get(
        f"/api/applications/{application_id}/comparison/workbook"
    )
    assert exported.status_code == 200
    disposition = exported.headers["content-disposition"]
    assert f'filename="priced-{application_id}.xlsx"' in disposition
    assert "filename*=UTF-8''" in disposition


def test_workbook_export_returns_retry_when_generation_is_busy(
    client: object,
    repository_root: Path,
) -> None:
    from app.data import workbook_export_slots

    detail = _upload(client, repository_root / "samples" / "synthetic_intake_completed.xlsx")
    application_id = detail["application"]["id"]
    _resolve_demo_gaps(client, application_id)
    assert workbook_export_slots.acquire(blocking=False)
    try:
        response = client.get(
            f"/api/applications/{application_id}/comparison/workbook"
        )
    finally:
        workbook_export_slots.release()
    assert response.status_code == 429
    assert response.headers["retry-after"] == "5"


def test_rhel_assumption_is_visible_on_priced_lines(
    client: object,
    repository_root: Path,
) -> None:
    detail = _upload(client, repository_root / "samples" / "synthetic_intake_completed.xlsx")
    completed = _resolve_demo_gaps(client, detail["application"]["id"])
    rhel_lines = [
        line
        for line in completed["comparison"]["line_items"]
        if line["component"] == "RHEL subscription uplift"
    ]
    assert rhel_lines
    assert all(line["demo_assumption"] for line in rhel_lines)
    assert all("visible demo assumption" in line["evidence"] for line in rhel_lines)


def test_workbook_export_is_locked_until_completeness_gate_passes(
    client: object,
    repository_root: Path,
) -> None:
    detail = _upload(client, repository_root / "samples" / "synthetic_intake_completed.xlsx")
    response = client.get(
        f"/api/applications/{detail['application']['id']}/comparison/workbook"
    )
    assert response.status_code == 409
    assert "CompletenessGate" in response.json()["detail"]


def test_workbook_retention_is_bounded_by_total_source_bytes(
    client: object,
    repository_root: Path,
    monkeypatch: object,
) -> None:
    source = repository_root / "samples" / "synthetic_intake_completed.xlsx"
    monkeypatch.setattr("app.data.MAX_STORED_WORKBOOK_BYTES", source.stat().st_size)
    _upload(client, source)

    with source.open("rb") as workbook:
        response = client.post(
            "/api/intakes",
            files={
                "file": (
                    source.name,
                    workbook,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
    assert response.status_code == 429


def test_missing_storage_performance_stays_unpriced(
    client: object,
    repository_root: Path,
    tmp_path: Path,
) -> None:
    source = repository_root / "samples" / "synthetic_intake_completed.xlsx"
    workbook = load_workbook(source)
    workbook["4 Storage"]["H2"] = ""
    modified_path = tmp_path / "missing-performance.xlsx"
    workbook.save(modified_path)

    detail = _upload(client, modified_path)
    application_id = detail["application"]["id"]
    gap_ids = {gap["id"] for gap in detail["gaps"]}
    assert "storage-1.performance" in gap_ids

    client.post(
        f"/api/applications/{application_id}/gaps/P1.regions/resolve",
        json={
            "resolved_by": "FinOps Reviewer",
            "azure_region": "eastus2",
            "aws_region": "us-east-1",
        },
    )
    current = client.get(f"/api/applications/{application_id}").json()["comparison"]
    assert current["state"] == "DraftBenchmark"
    assert any(
        line["skumap_component_id"] == "block-iops"
        and line["status"] == "Unpriced"
        for line in current["line_items"]
    )


def test_formula_cell_is_rejected(client: object, repository_root: Path, tmp_path: Path) -> None:
    source = repository_root / "samples" / "synthetic_intake_completed.xlsx"
    workbook = load_workbook(source)
    workbook["2 Servers"]["G2"] = "=2+2"
    modified_path = tmp_path / "formula.xlsx"
    workbook.save(modified_path)

    with modified_path.open("rb") as upload:
        response = client.post(
            "/api/intakes",
            files={"file": (modified_path.name, upload, "application/octet-stream")},
        )
    assert response.status_code == 422
    assert "formulas are not accepted" in response.json()["detail"]


def test_oversized_intake_request_is_rejected_before_parsing(client: object) -> None:
    response = client.post(
        "/api/intakes",
        headers={"Content-Length": str(20 * 1024 * 1024)},
        content=b"",
    )
    assert response.status_code == 413


def test_failed_resolution_is_transactional(
    client: object,
    repository_root: Path,
) -> None:
    from app.pricing import pricing_engine

    detail = _upload(client, repository_root / "samples" / "synthetic_intake_completed.xlsx")
    application_id = detail["application"]["id"]
    original_digest = pricing_engine._approval_digest
    pricing_engine._approval_digest = "tampered"
    try:
        response = client.post(
            f"/api/applications/{application_id}/gaps/P1.regions/resolve",
            json={
                "resolved_by": "Reviewer",
                "azure_region": "eastus2",
                "aws_region": "us-east-1",
            },
        )
        assert response.status_code == 503
    finally:
        pricing_engine._approval_digest = original_digest

    current = client.get(f"/api/applications/{application_id}").json()
    region_gap = next(gap for gap in current["gaps"] if gap["id"] == "P1.regions")
    assert region_gap["status"] == "Open"


def test_application_listing_uses_cached_comparisons(
    client: object,
    repository_root: Path,
    monkeypatch: object,
) -> None:
    _upload(client, repository_root / "samples" / "synthetic_intake_completed.xlsx")

    def fail_if_recalculated(*args: object, **kwargs: object) -> None:
        raise AssertionError("list/get must not reprice cached records")

    monkeypatch.setattr("app.data.pricing_engine.compare", fail_if_recalculated)
    assert client.get("/api/applications").status_code == 200
