import io
from copy import copy

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties

from app.models import ApplicationDetail, Cloud

SUMMARY_SHEET = "Pricing Summary"
DETAIL_SHEET = "Pricing Detail"
ASSUMPTIONS_SHEET = "Pricing Assumptions"
EXCLUSIONS_SHEET = "Pricing Exclusions"
EVIDENCE_SHEET = "Pricing Evidence"

NAVY = "17365D"
BLUE = "0000FF"
GREEN = "008000"
LIGHT_BLUE = "D9EAF7"
LIGHT_GRAY = "E7E6E6"
YELLOW = "FFF2CC"
WHITE = "FFFFFF"
BLACK = "000000"
CURRENCY_FORMAT = '$#,##0.00;($#,##0.00);-'
RATE_FORMAT = '$0.000000'
QUANTITY_FORMAT = '#,##0.0000'


class UnsafeSourceWorkbookError(ValueError):
    """The uploaded workbook carries a formula, so it can't be copied into a deliverable that recalculates."""


def build_priced_workbook(
    source: bytes,
    detail: ApplicationDetail,
    view: str | None = None,
) -> bytes:
    """Build the priced workbook. `view="aws"` selects the AWS-only layout; otherwise the
    comparison's presentation policy picks the Azure-only or both-cloud layout."""
    aws_only = view == "aws"
    workbook = load_workbook(io.BytesIO(source), data_only=False, keep_links=False)
    # Intake refuses formulas; this is the second line, since the export recalculates on open.
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    raise UnsafeSourceWorkbookError(
                        f"The uploaded workbook has a formula in {sheet.title}!{cell.coordinate}, "
                        "so it can't be exported."
                    )
        if len(sheet.conditional_formatting) or sheet.data_validations.dataValidation:
            raise UnsafeSourceWorkbookError(
                f"The uploaded workbook has conditional formatting or data validation on {sheet.title}, "
                "whose formulas Excel evaluates, so it can't be exported."
            )
        for table in sheet.tables.values():
            for column in table.tableColumns:
                if (
                    column.calculatedColumnFormula is not None
                    or column.totalsRowFormula is not None
                    or (column.totalsRowFunction or "none") != "none"
                ):
                    raise UnsafeSourceWorkbookError(
                        f"The uploaded workbook has a calculated table column on {sheet.title}, "
                        "so it can't be exported."
                    )
    if len(workbook.defined_names) or any(len(sheet.defined_names) for sheet in workbook.worksheets):
        raise UnsafeSourceWorkbookError(
            "The uploaded workbook has defined names, which can carry formulas, so it can't be exported."
        )
    names = {
        "summary": _unique_sheet_title(workbook, SUMMARY_SHEET),
        "detail": _unique_sheet_title(workbook, DETAIL_SHEET),
        "assumptions": _unique_sheet_title(workbook, ASSUMPTIONS_SHEET),
        "exclusions": _unique_sheet_title(workbook, EXCLUSIONS_SHEET),
        "evidence": _unique_sheet_title(workbook, EVIDENCE_SHEET),
    }
    summary = workbook.create_sheet(names["summary"], 0)
    pricing_detail = workbook.create_sheet(names["detail"], 1)
    assumptions = workbook.create_sheet(names["assumptions"], 2)
    exclusions = workbook.create_sheet(names["exclusions"], 3)
    evidence = workbook.create_sheet(names["evidence"], 4)
    names.update(
        {
            "summary": summary.title,
            "detail": pricing_detail.title,
            "assumptions": assumptions.title,
            "exclusions": exclusions.title,
            "evidence": evidence.title,
        }
    )

    if aws_only:
        _build_aws_detail(pricing_detail, detail)
        _build_aws_summary(summary, detail, pricing_detail.max_row, names)
    else:
        _build_detail(pricing_detail, detail)
        _build_summary(summary, detail, pricing_detail.max_row, names)
    _build_assumptions(assumptions, detail, aws_only=aws_only)
    _build_exclusions(exclusions, detail)
    _build_evidence(evidence, detail, aws_only=aws_only)

    if workbook.calculation is None:
        workbook.calculation = CalcProperties()
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def _build_summary(
    sheet,
    detail: ApplicationDetail,
    detail_last_row: int,
    names: dict[str, str],
) -> None:
    _build_azure_summary(sheet, detail, detail_last_row, names)


def _build_azure_summary(
    sheet,
    detail: ApplicationDetail,
    detail_last_row: int,
    names: dict[str, str],
) -> None:
    comparison = detail.comparison
    scenarios = {
        (scenario.provider.value, scenario.offer.value): scenario
        for scenario in comparison.commercial_scenarios
    }
    show_aws = comparison.presentation.show_aws
    _title(
        sheet,
        "Cloud Pricing Summary" if show_aws else "Azure Pricing Summary",
        "A1:F1",
    )
    sheet["A3"] = "Application"
    sheet["B3"] = detail.application.name
    sheet["A4"] = "Benchmark status"
    sheet["B4"] = comparison.state.value
    sheet["A5"] = "Priced as of"
    sheet["B5"] = comparison.priced_as_of
    sheet["A6"] = "Presentation"
    sheet["B6"] = (
        "AWS and Azure public pricing scenarios"
        if show_aws else "Azure public pricing scenarios"
    )
    _header_row(
        sheet,
        8,
        [
            "Public scenario" if show_aws else "Azure scenario",
            "Workbook total",
            "Engine total",
            "Reconciliation",
            "Covered",
            "Uncovered",
        ],
    )
    detail_ref = _sheet_ref(names["detail"])
    rows = (
        (9, "Azure", "List", "Public List", "H", None, None),
        (10, "Azure", "SavingsPlan", "3-year Savings Plan", "K", "I", "J"),
        (11, "Azure", "Reservation", "3-year Reservation", "N", "L", "M"),
    )
    if show_aws:
        rows += (
            (12, "AWS", "List", "AWS public On-Demand", "S", None, None),
            (13, "AWS", "SavingsPlan", "AWS Compute Savings Plan", "V", "T", "U"),
            (14, "AWS", "Reservation", "AWS Standard Reserved Instance", "Y", "W", "X"),
        )
    for (
        row, provider, offer, label, total_column, covered_column, uncovered_column
    ) in rows:
        scenario = scenarios[(provider, offer)]
        sheet.cell(
            row=row,
            column=1,
            value=scenario.label if show_aws else label,
        )
        if scenario.available:
            sheet.cell(
                row=row,
                column=2,
                value=f"=ROUND(SUM({detail_ref}!{total_column}6:{total_column}{detail_last_row}),2)",
            )
            sheet.cell(row=row, column=3, value=scenario.monthly_total)
            basis_check = (
                {"I": "AH", "L": "AI", "T": "AJ", "W": "AK"}.get(covered_column)
                if show_aws else None
            )
            sheet.cell(
                row=row,
                column=4,
                value=(
                    f'=IF(OR(ABS(B{row}-C{row})>=0.005,'
                    f'COUNTIF({detail_ref}!{basis_check}6:{basis_check}{detail_last_row},'
                    f'"MISMATCH")>0),"MISMATCH","Reconciled")'
                    if basis_check else
                    f'=IF(ABS(B{row}-C{row})<0.005,"Reconciled","MISMATCH")'
                ),
            )
            if covered_column is None:
                sheet.cell(row=row, column=5, value=0)
                sheet.cell(row=row, column=6, value=f"=B{row}")
            else:
                sheet.cell(
                    row=row,
                    column=5,
                    value=f"=ROUND(SUM({detail_ref}!{covered_column}6:{covered_column}{detail_last_row}),2)",
                )
                sheet.cell(
                    row=row,
                    column=6,
                    value=f"=ROUND(SUM({detail_ref}!{uncovered_column}6:{uncovered_column}{detail_last_row}),2)",
                )
        else:
            sheet.cell(row=row, column=2, value="Unavailable")
            sheet.cell(row=row, column=4, value=scenario.unavailable_reason)
    for row in range(9, 15 if show_aws else 12):
        for column in (2, 5, 6):
            sheet.cell(row=row, column=column).number_format = CURRENCY_FORMAT
            sheet.cell(row=row, column=column).font = Font(
                name="Arial", color=GREEN
            )
        sheet.cell(row=row, column=3).number_format = CURRENCY_FORMAT
        sheet.cell(row=row, column=3).font = Font(name="Arial", color=BLUE)

    _header_row(
        sheet,
        17 if show_aws else 14,
        [
            "Public basis",
            "Reference discount",
            "Target discount to parity" if show_aws else "Azure discount to parity",
            "Additional target advantage" if show_aws else "Additional Azure advantage",
            "Disclosure",
            "Target provider" if show_aws else "",
        ],
    )
    row = 18 if show_aws else 15
    for sensitivity in comparison.breakeven_sensitivities:
        if not sensitivity.available:
            sheet.cell(row=row, column=1, value=sensitivity.label)
            sheet.cell(row=row, column=2, value="Unavailable")
            sheet.cell(
                row=row,
                column=3,
                value="Complete public-price basis required.",
            )
            sheet.cell(
                row=row,
                column=5,
                value=(
                    "Hypothetical sensitivity remains unavailable until the "
                    "approved public-price basis is complete."
                ),
            )
            sheet.cell(row=row, column=5).alignment = Alignment(wrap_text=True)
            row += 1
            continue
        for point in sensitivity.points:
            sheet.cell(row=row, column=1, value=sensitivity.label)
            sheet.cell(
                row=row,
                column=2,
                value=point.reference_discount_percent / 100,
            )
            if show_aws:
                sheet.cell(
                    row=row,
                    column=6,
                    value=(
                        sensitivity.target_provider.value
                        if sensitivity.target_provider
                        else "Equal"
                    ),
                )
            if show_aws or (
                sensitivity.target_provider
                and sensitivity.target_provider.value == "Azure"
            ):
                sheet.cell(
                    row=row,
                    column=3,
                    value=point.target_discount_to_parity_percent / 100,
                )
                sheet.cell(
                    row=row,
                    column=4,
                    value=point.additional_discount_advantage_percent / 100,
                )
            else:
                sheet.cell(row=row, column=3, value="Already at or below parity")
                sheet.cell(row=row, column=4, value=0)
            sheet.cell(
                row=row,
                column=5,
                value=(
                    "Azure RHEL demo assumptions affect modeled parity; not a provider verdict. "
                    if show_aws and any(line.demo_assumption for line in comparison.line_items)
                    else ""
                ) + sensitivity.disclosure,
            )
            for column in (2, 3, 4):
                sheet.cell(row=row, column=column).number_format = "0.0%"
                sheet.cell(row=row, column=column).font = Font(
                    name="Arial", color=BLUE
                )
            sheet.cell(row=row, column=5).alignment = Alignment(wrap_text=True)
            row += 1

    notes_row = row + 1
    _header_row(sheet, notes_row, ["Important notes", "", "", "", "", ""])
    notes = [
        "This workbook is an estimate, not a purchasing decision or contracted quote.",
        *[
            warning if show_aws else _azure_presentation_warning(warning)
            for warning in detail.comparison.warnings
        ],
        f"Material omissions remain visible on the {names['exclusions']} sheet.",
        "The evidence JSON retains the complete calculation record.",
    ]
    for note_row, note in enumerate(notes, start=notes_row + 1):
        sheet.cell(row=note_row, column=1, value=note)
        sheet.merge_cells(
            start_row=note_row,
            start_column=1,
            end_row=note_row,
            end_column=6,
        )
        sheet.cell(row=note_row, column=1).alignment = Alignment(wrap_text=True)
    _format_sheet(sheet, {1: 29, 2: 20, 3: 22, 4: 24, 5: 72, 6: 18})


def _build_detail(sheet, detail: ApplicationDetail) -> None:
    _build_azure_detail(sheet, detail)


def _build_azure_detail(sheet, detail: ApplicationDetail) -> None:
    show_aws = detail.comparison.presentation.show_aws
    _title(
        sheet,
        "Cloud Pricing Detail" if show_aws else "Azure Pricing Detail",
        "A1:AK1" if show_aws else "A1:O1",
    )
    sheet["A3"] = (
        "Blue values are approved public-price inputs and line amounts; "
        "yellow highlights demo assumptions; black cells are workbook calculations. "
        "Commitment basis rates and quantities at Z:AG are verified against server-computed covered amounts."
        if show_aws
        else (
            "Blue values are approved public-price inputs; black cells are workbook "
            "calculations. Covered and uncovered commitment amounts remain explicit."
        )
    )
    sheet.merge_cells("A3:AK3" if show_aws else "A3:O3")
    sheet["A3"].fill = PatternFill("solid", fgColor=YELLOW)
    sheet["A3"].alignment = Alignment(wrap_text=True)
    headers = [
        "Unit",
        "Component",
        "Status",
        "Match",
        "Azure quantity",
        "Quantity unit",
        "List rate",
        "List monthly",
        "Savings Plan covered",
        "Savings Plan uncovered",
        "Savings Plan monthly",
        "Reservation covered",
        "Reservation uncovered",
        "Reservation monthly",
        "Evidence / reason",
    ]
    if show_aws:
        for column in (7, 8, 9, 10, 11, 12, 13, 14):
            headers[column - 1] = f"Azure {headers[column - 1]}"
        headers.extend(
            [
                "AWS quantity",
                "AWS quantity unit",
                "AWS List rate",
                "AWS List monthly",
                "AWS Savings Plan covered",
                "AWS Savings Plan uncovered",
                "AWS Savings Plan monthly",
                "AWS RI covered",
                "AWS RI uncovered",
                "AWS RI monthly",
                "Azure Savings Plan rate",
                "Azure Savings Plan billed quantity",
                "Azure Reservation rate",
                "Azure Reservation billed quantity",
                "AWS Savings Plan rate",
                "AWS Savings Plan billed quantity",
                "AWS RI rate",
                "AWS RI billed quantity",
                "Azure Savings Plan basis check",
                "Azure Reservation basis check",
                "AWS Savings Plan basis check",
                "AWS RI basis check",
            ]
        )
    _header_row(sheet, 5, headers)
    for row, line in enumerate(detail.comparison.line_items, start=6):
        values = [
            line.unit_name,
            line.component,
            line.status.value,
            line.match_class.value,
            line.azure_quantity,
            _azure_quantity_unit(line.quantity_unit),
            line.azure_commercial.list_rate,
        ]
        for column, value in enumerate(values, start=1):
            sheet.cell(row=row, column=column, value=value)
        sheet.cell(
            row=row,
            column=8,
            value=(
                f'=IF(OR(C{row}<>"Priced",E{row}="",G{row}=""),"",'
                f"ROUND(E{row}*G{row},2))"
            ),
        )
        sheet.cell(
            row=row,
            column=9,
            value=line.azure_commercial.savings_plan_covered_amount,
        )
        sheet.cell(
            row=row,
            column=10,
            value=line.azure_commercial.savings_plan_uncovered_amount,
        )
        sheet.cell(
            row=row,
            column=11,
            value=(
                f'=IF(COUNT(I{row}:J{row})=2,ROUND(SUM(I{row}:J{row}),2),"")'
                if show_aws
                else f'=IF(COUNT(I{row}:J{row})=0,"",ROUND(SUM(I{row}:J{row}),2))'
            ),
        )
        sheet.cell(
            row=row,
            column=12,
            value=line.azure_commercial.reservation_covered_amount,
        )
        sheet.cell(
            row=row,
            column=13,
            value=line.azure_commercial.reservation_uncovered_amount,
        )
        sheet.cell(
            row=row,
            column=14,
            value=(
                f'=IF(COUNT(L{row}:M{row})=2,ROUND(SUM(L{row}:M{row}),2),"")'
                if show_aws
                else f'=IF(COUNT(L{row}:M{row})=0,"",ROUND(SUM(L{row}:M{row}),2))'
            ),
        )
        sheet.cell(
            row=row,
            column=15,
            value=(
                line.unpriced_reason
                or (line.evidence if show_aws else None)
                or (
                    "Azure demo assumption; see Pricing Evidence."
                    if line.demo_assumption
                    else "Approved Azure PriceBook meter; see Pricing Evidence."
                )
            ),
        )
        sheet.cell(row=row, column=5).number_format = QUANTITY_FORMAT
        sheet.cell(row=row, column=7).number_format = RATE_FORMAT
        for column in (7, 9, 10, 12, 13):
            sheet.cell(row=row, column=column).font = Font(
                name="Arial", color=BLUE
            )
        for column in (8, 9, 10, 11, 12, 13, 14):
            sheet.cell(row=row, column=column).number_format = CURRENCY_FORMAT
        sheet.cell(row=row, column=15).alignment = Alignment(wrap_text=True)
        if show_aws:
            aws = line.aws_commercial
            for column, value in (
                (16, line.aws_quantity),
                (17, _aws_quantity_unit(line.quantity_unit)),
                (18, aws.list_rate),
                (
                    19,
                    f'=IF(OR(C{row}<>"Priced",P{row}="",R{row}=""),"",'
                    f"ROUND(P{row}*R{row},2))",
                ),
                (20, aws.savings_plan_covered_amount),
                (21, aws.savings_plan_uncovered_amount),
                (22, f'=IF(COUNT(T{row}:U{row})=2,ROUND(SUM(T{row}:U{row}),2),"")'),
                (23, aws.reservation_covered_amount),
                (24, aws.reservation_uncovered_amount),
                (25, f'=IF(COUNT(W{row}:X{row})=2,ROUND(SUM(W{row}:X{row}),2),"")'),
            ):
                sheet.cell(row=row, column=column, value=value)
            sheet.cell(row=row, column=16).number_format = QUANTITY_FORMAT
            sheet.cell(row=row, column=18).number_format = RATE_FORMAT
            for column in (19, 20, 21, 22, 23, 24, 25):
                sheet.cell(row=row, column=column).number_format = CURRENCY_FORMAT
            for column in (8, 18, 19, 20, 21, 23, 24):
                sheet.cell(row=row, column=column).font = Font(
                    name="Arial", color=BLUE
                )
            for rate_column, quantity_column, check_column, covered_column, rate, quantity, covered in (
                (26, 27, 34, "I", line.azure_commercial.savings_plan_rate,
                 line.azure_commercial.savings_plan_quantity,
                 line.azure_commercial.savings_plan_covered_amount),
                (28, 29, 35, "L", line.azure_commercial.reservation_rate,
                 line.azure_commercial.reservation_quantity,
                 line.azure_commercial.reservation_covered_amount),
                (30, 31, 36, "T", aws.savings_plan_rate,
                 aws.savings_plan_quantity, aws.savings_plan_covered_amount),
                (32, 33, 37, "W", aws.reservation_rate,
                 aws.reservation_quantity, aws.reservation_covered_amount),
            ):
                if covered is not None and covered > 0:
                    if rate is None or quantity is None:
                        raise ValueError("Covered commitment lacks a rate or billed quantity.")
                    sheet.cell(row=row, column=rate_column, value=rate)
                    sheet.cell(row=row, column=quantity_column, value=quantity)
                    rate_ref = sheet.cell(row=row, column=rate_column).column_letter
                    quantity_ref = sheet.cell(row=row, column=quantity_column).column_letter
                    sheet.cell(
                        row=row, column=check_column,
                        value=(
                            f'=IF(ABS(ROUND({rate_ref}{row}*{quantity_ref}{row},2)'
                            f'-{covered_column}{row})<0.005,"Reconciled","MISMATCH")'
                        ),
                    )
                    sheet.cell(row=row, column=rate_column).number_format = RATE_FORMAT
                    sheet.cell(row=row, column=quantity_column).number_format = QUANTITY_FORMAT
                    sheet.cell(row=row, column=rate_column).font = Font(name="Arial", color=BLUE)
                    sheet.cell(row=row, column=quantity_column).font = Font(name="Arial", color=BLUE)
            if line.demo_assumption:
                for column in (7, 8, 10, 11, 13, 14):
                    sheet.cell(row=row, column=column).fill = PatternFill(
                        "solid", fgColor=YELLOW
                    )
    sheet.freeze_panes = "A6"
    sheet.auto_filter.ref = f"A5:{'AK' if show_aws else 'O'}{sheet.max_row}"
    widths = {
        1: 22,
        2: 30,
        3: 13,
        4: 20,
        5: 16,
        6: 18,
        7: 14,
        8: 16,
        9: 19,
        10: 19,
        11: 18,
        12: 19,
        13: 19,
        14: 18,
        15: 58,
    }
    if show_aws:
        widths.update(
            {
                16: 17,
                17: 24,
                18: 17,
                19: 19,
                20: 23,
                21: 23,
                22: 23,
                23: 20,
                24: 20,
                25: 20,
            }
        )
        widths.update({column: 23 for column in range(26, 38)})
    _format_sheet(sheet, widths)


def _aws_placeholder_line_ids(detail: ApplicationDetail) -> set[str]:
    return {
        line_id
        for sensitivity in detail.comparison.assumption_sensitivities
        if sensitivity.provider == Cloud.AWS
        for line_id in sensitivity.line_ids
    }


def _build_aws_summary(
    sheet,
    detail: ApplicationDetail,
    detail_last_row: int,
    names: dict[str, str],
) -> None:
    # Mirrors the Azure-only layout cell for cell, with AWS amounts, so neither cloud reads differently.
    comparison = detail.comparison
    scenarios = {
        (scenario.provider.value, scenario.offer.value): scenario
        for scenario in comparison.commercial_scenarios
    }
    _title(sheet, "AWS Pricing Summary", "A1:F1")
    sheet["A3"] = "Application"
    sheet["B3"] = detail.application.name
    sheet["A4"] = "Benchmark status"
    sheet["B4"] = comparison.state.value
    sheet["A5"] = "Priced as of"
    sheet["B5"] = comparison.priced_as_of
    sheet["A6"] = "Presentation"
    sheet["B6"] = "AWS public pricing scenarios"
    _header_row(
        sheet,
        8,
        ["AWS scenario", "Workbook total", "Engine total", "Reconciliation", "Covered", "Uncovered"],
    )
    detail_ref = _sheet_ref(names["detail"])
    rows = (
        (9, "List", "Public On-Demand", "H", None, None),
        (10, "SavingsPlan", "3-year Compute Savings Plan", "K", "I", "J"),
        (11, "Reservation", "3-year Standard Reserved Instance", "N", "L", "M"),
    )
    for row, offer, label, total_column, covered_column, uncovered_column in rows:
        scenario = scenarios[("AWS", offer)]
        sheet.cell(row=row, column=1, value=label)
        if scenario.available:
            sheet.cell(
                row=row,
                column=2,
                value=f"=ROUND(SUM({detail_ref}!{total_column}6:{total_column}{detail_last_row}),2)",
            )
            sheet.cell(row=row, column=3, value=scenario.monthly_total)
            sheet.cell(
                row=row,
                column=4,
                value=f'=IF(ABS(B{row}-C{row})<0.005,"Reconciled","MISMATCH")',
            )
            if covered_column is None:
                sheet.cell(row=row, column=5, value=0)
                sheet.cell(row=row, column=6, value=f"=B{row}")
            else:
                sheet.cell(
                    row=row,
                    column=5,
                    value=f"=ROUND(SUM({detail_ref}!{covered_column}6:{covered_column}{detail_last_row}),2)",
                )
                sheet.cell(
                    row=row,
                    column=6,
                    value=f"=ROUND(SUM({detail_ref}!{uncovered_column}6:{uncovered_column}{detail_last_row}),2)",
                )
        else:
            sheet.cell(row=row, column=2, value="Unavailable")
            sheet.cell(row=row, column=4, value=scenario.unavailable_reason)
    for row in range(9, 12):
        for column in (2, 5, 6):
            sheet.cell(row=row, column=column).number_format = CURRENCY_FORMAT
            sheet.cell(row=row, column=column).font = Font(name="Arial", color=GREEN)
        sheet.cell(row=row, column=3).number_format = CURRENCY_FORMAT
        sheet.cell(row=row, column=3).font = Font(name="Arial", color=BLUE)

    _header_row(
        sheet,
        14,
        [
            "Public basis",
            "Reference discount",
            "AWS discount to parity",
            "Additional AWS advantage",
            "Disclosure",
            "",
        ],
    )
    row = 15
    for sensitivity in comparison.breakeven_sensitivities:
        if not sensitivity.available:
            sheet.cell(row=row, column=1, value=sensitivity.label)
            sheet.cell(row=row, column=2, value="Unavailable")
            sheet.cell(row=row, column=3, value="Complete public-price basis required.")
            sheet.cell(
                row=row,
                column=5,
                value=(
                    "Hypothetical sensitivity remains unavailable until the "
                    "approved public-price basis is complete."
                ),
            )
            sheet.cell(row=row, column=5).alignment = Alignment(wrap_text=True)
            row += 1
            continue
        for point in sensitivity.points:
            sheet.cell(row=row, column=1, value=sensitivity.label)
            sheet.cell(row=row, column=2, value=point.reference_discount_percent / 100)
            if sensitivity.target_provider == Cloud.AWS:
                sheet.cell(row=row, column=3, value=point.target_discount_to_parity_percent / 100)
                sheet.cell(
                    row=row,
                    column=4,
                    value=point.additional_discount_advantage_percent / 100,
                )
            else:
                sheet.cell(row=row, column=3, value="Already at or below parity")
                sheet.cell(row=row, column=4, value=0)
            sheet.cell(row=row, column=5, value=sensitivity.disclosure)
            for column in (2, 3, 4):
                sheet.cell(row=row, column=column).number_format = "0.0%"
                sheet.cell(row=row, column=column).font = Font(name="Arial", color=BLUE)
            sheet.cell(row=row, column=5).alignment = Alignment(wrap_text=True)
            row += 1

    notes_row = row + 1
    _header_row(sheet, notes_row, ["Important notes", "", "", "", "", ""])
    notes = [
        "This workbook is an estimate, not a purchasing decision or contracted quote.",
        *[_aws_presentation_warning(warning) for warning in comparison.warnings],
        f"Material omissions remain visible on the {names['exclusions']} sheet.",
        "The evidence JSON retains the complete calculation record.",
    ]
    for note_row, note in enumerate(notes, start=notes_row + 1):
        sheet.cell(row=note_row, column=1, value=note)
        sheet.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=6)
        sheet.cell(row=note_row, column=1).alignment = Alignment(wrap_text=True)
    _format_sheet(sheet, {1: 29, 2: 20, 3: 22, 4: 24, 5: 72, 6: 18})


def _build_aws_detail(sheet, detail: ApplicationDetail) -> None:
    placeholder_lines = _aws_placeholder_line_ids(detail)
    _title(sheet, "AWS Pricing Detail", "A1:O1")
    sheet["A3"] = (
        "Blue values are approved public-price inputs; black cells are workbook "
        "calculations. Covered and uncovered commitment amounts remain explicit."
    )
    sheet.merge_cells("A3:O3")
    sheet["A3"].fill = PatternFill("solid", fgColor=YELLOW)
    sheet["A3"].alignment = Alignment(wrap_text=True)
    _header_row(
        sheet,
        5,
        [
            "Unit",
            "Component",
            "Status",
            "Match",
            "AWS quantity",
            "Quantity unit",
            "List rate",
            "List monthly",
            "Savings Plan covered",
            "Savings Plan uncovered",
            "Savings Plan monthly",
            "Reserved Instance covered",
            "Reserved Instance uncovered",
            "Reserved Instance monthly",
            "Evidence / reason",
        ],
    )
    for row, line in enumerate(detail.comparison.line_items, start=6):
        aws = line.aws_commercial
        placeholder = line.id in placeholder_lines
        values = [
            line.unit_name,
            line.component,
            line.status.value,
            line.match_class.value,
            line.aws_quantity,
            _aws_quantity_unit(line.quantity_unit),
            aws.list_rate,
        ]
        for column, value in enumerate(values, start=1):
            sheet.cell(row=row, column=column, value=value)
        sheet.cell(
            row=row,
            column=8,
            value=(
                f'=IF(OR(C{row}<>"Priced",E{row}="",G{row}=""),"",'
                f"ROUND(E{row}*G{row},2))"
            ),
        )
        sheet.cell(row=row, column=9, value=aws.savings_plan_covered_amount)
        sheet.cell(row=row, column=10, value=aws.savings_plan_uncovered_amount)
        sheet.cell(
            row=row,
            column=11,
            value=f'=IF(COUNT(I{row}:J{row})=0,"",ROUND(SUM(I{row}:J{row}),2))',
        )
        sheet.cell(row=row, column=12, value=aws.reservation_covered_amount)
        sheet.cell(row=row, column=13, value=aws.reservation_uncovered_amount)
        sheet.cell(
            row=row,
            column=14,
            value=f'=IF(COUNT(L{row}:M{row})=0,"",ROUND(SUM(L{row}:M{row}),2))',
        )
        sheet.cell(
            row=row,
            column=15,
            value=(
                line.unpriced_reason
                or (
                    "AWS demo assumption; see Pricing Evidence."
                    if placeholder
                    else "Approved AWS PriceBook meter; see Pricing Evidence."
                )
            ),
        )
        sheet.cell(row=row, column=5).number_format = QUANTITY_FORMAT
        sheet.cell(row=row, column=7).number_format = RATE_FORMAT
        for column in (7, 9, 10, 12, 13):
            sheet.cell(row=row, column=column).font = Font(name="Arial", color=BLUE)
        for column in (8, 9, 10, 11, 12, 13, 14):
            sheet.cell(row=row, column=column).number_format = CURRENCY_FORMAT
        sheet.cell(row=row, column=15).alignment = Alignment(wrap_text=True)
        if placeholder:
            for column in (7, 8, 10, 11, 13, 14):
                sheet.cell(row=row, column=column).fill = PatternFill("solid", fgColor=YELLOW)
    sheet.freeze_panes = "A6"
    sheet.auto_filter.ref = f"A5:O{sheet.max_row}"
    _format_sheet(
        sheet,
        {1: 22, 2: 30, 3: 13, 4: 20, 5: 16, 6: 18, 7: 14, 8: 16, 9: 19, 10: 19,
         11: 18, 12: 19, 13: 19, 14: 18, 15: 58},
    )


def _build_assumptions(sheet, detail: ApplicationDetail, *, aws_only: bool = False) -> None:
    if aws_only:
        _build_aws_assumptions(sheet, detail)
        return
    _title(sheet, "Pricing Assumptions", "A1:D1")
    _header_row(sheet, 3, ["Assumption", "Value", "Source"])
    rows = [
        ("Currency", detail.comparison.assumptions.currency, "Approved calculation policy"),
        (
            "Estimate horizon (months)",
            detail.comparison.assumptions.horizon_months,
            "Approved calculation policy",
        ),
        (
            "Commercial view",
            detail.comparison.assumptions.commercial_view.value,
            "Approved calculation policy",
        ),
        (
            "Commitment term",
            f"{detail.comparison.assumptions.commitment_term_years} years",
            "Approved calculation policy",
        ),
        (
            "Commitment payment option",
            detail.comparison.assumptions.commitment_payment_option,
            "Approved calculation policy",
        ),
        (
            "Commitment utilization",
            f"{detail.comparison.assumptions.commitment_utilization_percent}%",
            "Approved calculation policy",
        ),
        (
            "Commitment tenancy",
            detail.comparison.assumptions.commitment_tenancy,
            "Approved calculation policy",
        ),
        (
            "Azure region",
            detail.comparison.assumptions.azure_region,
            "Resolved Intake assumption",
        ),
        (
            "PriceBook snapshot",
            detail.comparison.assumptions.pricebook_snapshot_id,
            "Published PriceBook manifest",
        ),
    ]
    if detail.comparison.presentation.show_aws:
        rows.insert(
            -1,
            (
                "AWS region",
                detail.comparison.assumptions.aws_region,
                "Resolved Intake assumption",
            ),
        )
    for row_number, row in enumerate(rows, start=4):
        for column, value in enumerate(row, start=1):
            sheet.cell(row=row_number, column=column, value=value)
        sheet.cell(row=row_number, column=2).font = Font(name="Arial", color=BLUE)
    next_row = 4 + len(rows) + 2
    headers = ["License assessment", "Azure treatment"]
    if detail.comparison.presentation.show_aws:
        headers.append("AWS treatment")
    headers.append("Evidence")
    _header_row(sheet, next_row, headers)
    for row_number, assessment in enumerate(
        detail.comparison.license_assessments,
        start=next_row + 1,
    ):
        sheet.cell(
            row=row_number,
            column=1,
            value=f"{assessment.unit_id}: {assessment.product}",
        )
        sheet.cell(row=row_number, column=2, value=assessment.azure_treatment)
        evidence_column = 3
        if detail.comparison.presentation.show_aws:
            sheet.cell(row=row_number, column=3, value=assessment.aws_treatment)
            evidence_column = 4
        sheet.cell(
            row=row_number,
            column=evidence_column,
            value="\n".join(
                assessment.evidence_urls
                if detail.comparison.presentation.show_aws
                else _azure_source_urls(assessment.evidence_urls)
            ),
        )
        for column in range(2, evidence_column + 1):
            sheet.cell(row=row_number, column=column).alignment = Alignment(
                wrap_text=True
            )
    _format_sheet(sheet, {1: 34, 2: 34, 3: 48, 4: 72})


def _build_aws_assumptions(sheet, detail: ApplicationDetail) -> None:
    assumptions = detail.comparison.assumptions
    _title(sheet, "Pricing Assumptions", "A1:D1")
    _header_row(sheet, 3, ["Assumption", "Value", "Source"])
    rows = [
        ("Currency", assumptions.currency, "Approved calculation policy"),
        ("Estimate horizon (months)", assumptions.horizon_months, "Approved calculation policy"),
        ("Commercial view", assumptions.commercial_view.value, "Approved calculation policy"),
        ("Commitment term", f"{assumptions.commitment_term_years} years", "Approved calculation policy"),
        ("Commitment payment option", assumptions.commitment_payment_option, "Approved calculation policy"),
        (
            "Commitment utilization",
            f"{assumptions.commitment_utilization_percent}%",
            "Approved calculation policy",
        ),
        ("Commitment tenancy", assumptions.commitment_tenancy, "Approved calculation policy"),
        ("AWS region", assumptions.aws_region, "Resolved Intake assumption"),
        ("PriceBook snapshot", assumptions.pricebook_snapshot_id, "Published PriceBook manifest"),
    ]
    for row_number, row in enumerate(rows, start=4):
        for column, value in enumerate(row, start=1):
            sheet.cell(row=row_number, column=column, value=value)
        sheet.cell(row=row_number, column=2).font = Font(name="Arial", color=BLUE)
    next_row = 4 + len(rows) + 2
    _header_row(sheet, next_row, ["License assessment", "AWS treatment", "Evidence"])
    for row_number, assessment in enumerate(
        detail.comparison.license_assessments,
        start=next_row + 1,
    ):
        sheet.cell(row=row_number, column=1, value=f"{assessment.unit_id}: {assessment.product}")
        sheet.cell(row=row_number, column=2, value=assessment.aws_treatment)
        sheet.cell(
            row=row_number,
            column=3,
            value="\n".join(_aws_source_urls(assessment.evidence_urls)),
        )
        for column in (2, 3):
            sheet.cell(row=row_number, column=column).alignment = Alignment(wrap_text=True)
    _format_sheet(sheet, {1: 34, 2: 34, 3: 48, 4: 72})


def _build_exclusions(sheet, detail: ApplicationDetail) -> None:
    _title(sheet, "Excluded Cost Ledger", "A1:D1")
    sheet["A2"] = (
        "These material categories are not included in the displayed provider totals."
    )
    sheet.merge_cells("A2:D2")
    sheet["A2"].fill = PatternFill("solid", fgColor=YELLOW)
    _header_row(sheet, 4, ["Category", "Materiality", "Rationale", "Policy ID"])
    for row_number, exclusion in enumerate(
        detail.comparison.excluded_cost_ledger,
        start=5,
    ):
        values = [
            exclusion.category,
            exclusion.materiality,
            exclusion.rationale,
            exclusion.policy_id,
        ]
        for column, value in enumerate(values, start=1):
            sheet.cell(row=row_number, column=column, value=value)
        sheet.cell(row=row_number, column=3).alignment = Alignment(wrap_text=True)
    sheet.freeze_panes = "A5"
    sheet.auto_filter.ref = f"A4:D{sheet.max_row}"
    _format_sheet(sheet, {1: 40, 2: 16, 3: 80, 4: 42})


def _build_evidence(sheet, detail: ApplicationDetail, *, aws_only: bool = False) -> None:
    comparison = detail.comparison
    show_all = comparison.presentation.show_aws and not aws_only
    _title(sheet, "Pricing Evidence", "A1:B1")
    _header_row(sheet, 3, ["Evidence", "Value"])
    rows = [
        ("Demo extract ID", comparison.pricebook_snapshot_id),
        ("Demo extract content hash", comparison.pricebook_content_hash),
        ("Source published snapshot", comparison.source_pricebook_snapshot_id),
        ("Source published content hash", comparison.source_pricebook_content_hash),
        (
            "Source catalogs",
            "\n".join(
                _aws_source_urls(comparison.source_urls)
                if aws_only
                else comparison.source_urls
                if show_all
                else _azure_source_urls(comparison.source_urls)
            ),
        ),
        ("Priced as of", comparison.priced_as_of),
        ("SkuMap digest", comparison.skumap_content_digest),
        ("SkuMap approver", comparison.skumap_approval.approver),
        ("SkuMap approved at", comparison.skumap_approval.approved_at.isoformat()),
        ("SkuMap non-production", comparison.skumap_approval.non_production),
        ("Calculator rule", comparison.calculator_rule_version),
        ("Run hash", comparison.run_hash),
    ]
    for row_number, row in enumerate(rows, start=4):
        for column, value in enumerate(row, start=1):
            sheet.cell(row=row_number, column=column, value=value)
        sheet.cell(row=row_number, column=2).font = Font(name="Arial", color=BLUE)

    source_header_row = 4 + len(rows) + 2
    _header_row(
        sheet,
        source_header_row,
        ["Rate key", "Type", "Provider", "SKU / meter", "Source row ID", "Note"],
    )
    if aws_only:
        rate_sources = [
            source
            for source in comparison.rate_sources
            if source.provider.casefold() != "azure"
        ]
    elif show_all:
        rate_sources = comparison.rate_sources
    else:
        rate_sources = [
            source
            for source in comparison.rate_sources
            if source.provider.casefold() != "aws"
        ]
    for row_number, source in enumerate(
        rate_sources,
        start=source_header_row + 1,
    ):
        values = [
            source.rate_key,
            source.source_type,
            source.provider,
            " / ".join(value for value in (source.sku, source.meter) if value),
            ", ".join(source.row_ids) if source.row_ids else source.row_id,
            source.note,
        ]
        for column, value in enumerate(values, start=1):
            sheet.cell(row=row_number, column=column, value=value)
        if source.source_type == "DemoAssumption":
            for column in range(1, 7):
                sheet.cell(row=row_number, column=column).fill = PatternFill(
                    "solid",
                    fgColor=YELLOW,
                )
        sheet.cell(row=row_number, column=6).alignment = Alignment(wrap_text=True)
    sheet.freeze_panes = f"A{source_header_row + 1}"
    sheet.auto_filter.ref = f"A{source_header_row}:F{sheet.max_row}"
    _format_sheet(sheet, {1: 42, 2: 22, 3: 12, 4: 52, 5: 72, 6: 70})


def _azure_source_urls(urls: list[str]) -> list[str]:
    return [
        url
        for url in urls
        if "azure" in url.casefold() or "microsoft" in url.casefold()
    ]


def _aws_source_urls(urls: list[str]) -> list[str]:
    return [
        url
        for url in urls
        if "aws" in url.casefold() or "amazon" in url.casefold()
    ]


def _azure_quantity_unit(value: str | None) -> str | None:
    if value is None or ";" not in value:
        return value
    azure_value = next(
        (
            part.strip()
            for part in value.split(";")
            if part.strip().casefold().startswith("azure ")
        ),
        None,
    )
    return azure_value.removeprefix("Azure ") if azure_value else value


def _aws_quantity_unit(value: str | None) -> str | None:
    if value is None or ";" not in value:
        return value
    aws_value = next(
        (
            part.strip()
            for part in value.split(";")
            if part.strip().casefold().startswith("aws ")
        ),
        None,
    )
    return aws_value.removeprefix("AWS ") if aws_value else value


def _azure_presentation_warning(value: str) -> str:
    return (
        value.replace(
            "for approved East US 2 and us-east-1 mappings",
            "for the approved East US 2 mapping",
        )
        .replace("AWS", "secondary-cloud")
    )


def _aws_presentation_warning(value: str) -> str:
    return (
        value.replace(
            "for approved East US 2 and us-east-1 mappings",
            "for the approved us-east-1 mapping",
        )
        .replace("Azure", "secondary-cloud")
    )


def _title(sheet, text: str, merged_range: str) -> None:
    sheet.merge_cells(merged_range)
    cell = sheet.cell(row=1, column=1, value=text)
    cell.font = Font(name="Arial", bold=True, color=WHITE, size=16)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 27


def _header_row(sheet, row: int, values: list[str]) -> None:
    for column, value in enumerate(values, start=1):
        cell = sheet.cell(row=row, column=column, value=value)
        cell.font = Font(name="Arial", bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def _format_sheet(sheet, widths: dict[int, float]) -> None:
    for row in sheet.iter_rows():
        for cell in row:
            if cell.font.name != "Arial":
                cell.font = copy(cell.font)
                cell.font = Font(
                    name="Arial",
                    size=cell.font.size,
                    bold=cell.font.bold,
                    italic=cell.font.italic,
                    color=cell.font.color,
                )
            cell.alignment = copy(cell.alignment)
            cell.alignment = Alignment(
                horizontal=cell.alignment.horizontal,
                vertical=cell.alignment.vertical or "top",
                wrap_text=cell.alignment.wrap_text,
            )
    for column, width in widths.items():
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.sheet_view.showGridLines = False


def _unique_sheet_title(workbook, requested: str) -> str:
    existing = {title.casefold() for title in workbook.sheetnames}
    if requested.casefold() not in existing:
        return requested
    suffix = 2
    while True:
        marker = f" ({suffix})"
        candidate = f"{requested[: 31 - len(marker)]}{marker}"
        if candidate.casefold() not in existing:
            return candidate
        suffix += 1


def _sheet_ref(title: str) -> str:
    return f"'{title.replace(chr(39), chr(39) * 2)}'"
